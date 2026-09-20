# -*- coding: utf-8 -*-
import base64
import hashlib
import json
import logging
import os
import re
import sqlite3
import threading
import time
import urllib.parse
import uuid
from pathlib import Path

import requests
from fastapi import FastAPI, HTTPException, File, UploadFile
from fastapi.responses import FileResponse, Response, StreamingResponse
from pydantic import BaseModel

import feishu
import resume_pdf
import xhs

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
log = logging.getLogger("qiuzhao")

ROOT = Path(__file__).parent
DB_PATH = ROOT / "jobs.db"

STAGES = ["未投递", "已投递", "简历评估", "AI面试", "一面", "二面", "三面", "HR面", "Offer", "已挂"]
JOB_FIELDS = ["company", "title", "base", "requirements", "deadline", "applied_at", "stage", "notes"]
DEFAULT_PROMPT_NAME = "默认模板"


def _all_stages():
    try:
        custom = json.loads(_get_setting("custom_stages", "[]") or "[]")
        custom = [str(x) for x in custom if str(x).strip()]
    except Exception:
        custom = []
    return list(STAGES) + custom

MODEL_PRESETS = [
    {"name": "DeepSeek", "base": "https://api.deepseek.com", "model": "deepseek-chat"},
    {"name": "OpenCode Go · DeepSeek V4 Pro", "base": "https://opencode.ai/zen/go/v1", "model": "deepseek-v4-pro"},
    {"name": "OpenCode Go · GLM-5.3", "base": "https://opencode.ai/zen/go/v1", "model": "glm-5.3"},
    {"name": "OpenCode Go · Kimi K2.7", "base": "https://opencode.ai/zen/go/v1", "model": "kimi-k2.7-code"},
    {"name": "通义千问 Qwen", "base": "https://dashscope.aliyuncs.com/compatible-mode/v1", "model": "qwen-plus"},
    {"name": "智谱 GLM", "base": "https://open.bigmodel.cn/api/paas/v4", "model": "glm-4-plus"},
    {"name": "豆包 Doubao", "base": "https://ark.cn-beijing.volces.com/api/v3", "model": "doubao-1-5-pro-32k-250115"},
    {"name": "Moonshot Kimi", "base": "https://api.moonshot.cn/v1", "model": "moonshot-v1-8k"},
    {"name": "Ollama 本地", "base": "http://localhost:11434/v1", "model": "llama3"},
]


def _load_env():
    p = ROOT / ".env"
    if not p.exists():
        return
    for line in p.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        k, v = line.split("=", 1)
        os.environ.setdefault(k.strip(), v.strip())


_load_env()
ENV = os.environ.get

app = FastAPI(title="OfferCopilot", version="0.2.0")

_LLM_SESSION = uuid.uuid4().hex


def _conn():
    c = sqlite3.connect(os.environ.get("OC_DB") or str(DB_PATH))
    c.row_factory = sqlite3.Row
    return c


def init_db():
    c = _conn()
    try:
        with c:
            c.execute(
                "CREATE TABLE IF NOT EXISTS jobs ("
                "id INTEGER PRIMARY KEY AUTOINCREMENT,"
                "company TEXT DEFAULT '',"
                "title TEXT DEFAULT '',"
                "base TEXT DEFAULT '',"
                "requirements TEXT DEFAULT '',"
                "deadline TEXT DEFAULT '',"
                "applied_at TEXT DEFAULT '',"
                "stage TEXT DEFAULT '未投递',"
                "notes TEXT DEFAULT '',"
                "jd_text TEXT DEFAULT '',"
                "created_at TEXT DEFAULT '',"
                "match_score INTEGER DEFAULT 0,"
                "match_detail TEXT DEFAULT '',"
                "matched_at TEXT DEFAULT '')"
            )
            c.execute(
                "CREATE TABLE IF NOT EXISTS resumes ("
                "id INTEGER PRIMARY KEY AUTOINCREMENT,"
                "name TEXT DEFAULT '',"
                "content TEXT DEFAULT '',"
                "is_active INTEGER DEFAULT 0,"
                "created_at TEXT DEFAULT '')"
            )
            c.execute(
                "CREATE TABLE IF NOT EXISTS interviews ("
                "id INTEGER PRIMARY KEY AUTOINCREMENT,"
                "job_id INTEGER,"
                "title TEXT DEFAULT '',"
                "resume_text TEXT DEFAULT '',"
                "messages TEXT DEFAULT '',"
                "report TEXT DEFAULT '',"
                "status TEXT DEFAULT 'ing',"
                "overall INTEGER DEFAULT 0,"
                "created_at TEXT DEFAULT '',"
                "finished_at TEXT DEFAULT '')"
            )
            c.execute(
                "CREATE TABLE IF NOT EXISTS evalsets ("
                "id INTEGER PRIMARY KEY AUTOINCREMENT,"
                "name TEXT DEFAULT '',"
                "created_at TEXT DEFAULT '')"
            )
            c.execute(
                "CREATE TABLE IF NOT EXISTS evalcases ("
                "id INTEGER PRIMARY KEY AUTOINCREMENT,"
                "evalset_id INTEGER,"
                "jd_text TEXT DEFAULT '',"
                "company TEXT DEFAULT '',"
                "title TEXT DEFAULT '',"
                "base TEXT DEFAULT '',"
                "deadline TEXT DEFAULT '',"
                "created_at TEXT DEFAULT '')"
            )
            c.execute(
                "CREATE TABLE IF NOT EXISTS prompt_versions ("
                "id INTEGER PRIMARY KEY AUTOINCREMENT,"
                "name TEXT DEFAULT '',"
                "prompt TEXT DEFAULT '',"
                "created_at TEXT DEFAULT '')"
            )
            c.execute(
                "CREATE TABLE IF NOT EXISTS eval_runs ("
                "id INTEGER PRIMARY KEY AUTOINCREMENT,"
                "evalset_id INTEGER,"
                "version_id INTEGER,"
                "per_case TEXT DEFAULT '',"
                "summary TEXT DEFAULT '',"
                "created_at TEXT DEFAULT '')"
            )
            if c.execute("SELECT COUNT(*) FROM prompt_versions").fetchone()[0] == 0:
                c.execute(
                    "INSERT INTO prompt_versions (name,prompt,created_at) VALUES (?,?,?)",
                    (DEFAULT_PROMPT_NAME, PARSE_PROMPT, time.strftime("%Y-%m-%d %H:%M")),
                )
            c.execute(
                "CREATE TABLE IF NOT EXISTS settings ("
                "key TEXT PRIMARY KEY,"
                "value TEXT DEFAULT '')"
            )
            c.execute(
                "CREATE TABLE IF NOT EXISTS cache ("
                "key TEXT PRIMARY KEY,"
                "value TEXT DEFAULT '',"
                "created_at TEXT DEFAULT '')"
            )
            c.execute(
                "CREATE TABLE IF NOT EXISTS usage ("
                "id INTEGER PRIMARY KEY AUTOINCREMENT,"
                "kind TEXT DEFAULT '',"
                "in_chars INTEGER DEFAULT 0,"
                "out_chars INTEGER DEFAULT 0,"
                "latency REAL DEFAULT 0,"
                "model TEXT DEFAULT '',"
                "created_at TEXT DEFAULT '')"
            )
            c.execute(
                "CREATE TABLE IF NOT EXISTS reviews ("
                "id INTEGER PRIMARY KEY AUTOINCREMENT,"
                "company TEXT DEFAULT '',"
                "title TEXT DEFAULT '',"
                "round TEXT DEFAULT '',"
                "date TEXT DEFAULT '',"
                "questions TEXT DEFAULT '',"
                "self_rating INTEGER DEFAULT 0,"
                "notes TEXT DEFAULT '',"
                "ai_review TEXT DEFAULT '',"
                "created_at TEXT DEFAULT '')"
            )
            c.execute(
                "CREATE TABLE IF NOT EXISTS jings ("
                "id INTEGER PRIMARY KEY AUTOINCREMENT,"
                "company TEXT DEFAULT '',"
                "title TEXT DEFAULT '',"
                "source_url TEXT DEFAULT '',"
                "content TEXT DEFAULT '',"
                "questions TEXT DEFAULT '',"
                "summary TEXT DEFAULT '',"
                "is_fav INTEGER DEFAULT 0,"
                "folder_id INTEGER DEFAULT 0,"
                "created_at TEXT DEFAULT '')"
            )
            c.execute(
                "CREATE TABLE IF NOT EXISTS jing_folders ("
                "id INTEGER PRIMARY KEY AUTOINCREMENT,"
                "name TEXT DEFAULT '',"
                "created_at TEXT DEFAULT '')"
            )
            c.execute(
                "CREATE TABLE IF NOT EXISTS exp_items ("
                "id INTEGER PRIMARY KEY AUTOINCREMENT,"
                "type TEXT DEFAULT '实习',"
                "company TEXT DEFAULT '',"
                "title TEXT DEFAULT '',"
                "role TEXT DEFAULT '',"
                "time_range TEXT DEFAULT '',"
                "content TEXT DEFAULT '',"
                "created_at TEXT DEFAULT '')"
            )
            c.execute(
                "CREATE TABLE IF NOT EXISTS op_log ("
                "id INTEGER PRIMARY KEY AUTOINCREMENT,"
                "module TEXT DEFAULT '',"
                "level TEXT DEFAULT 'info',"
                "message TEXT DEFAULT '',"
                "created_at TEXT DEFAULT '')"
            )
            cols = {r["name"] for r in c.execute("PRAGMA table_info(jobs)").fetchall()}
            for col, ddl in [("match_score", "INTEGER DEFAULT 0"), ("match_detail", "TEXT DEFAULT ''"), ("matched_at", "TEXT DEFAULT ''")]:
                if col not in cols:
                    c.execute("ALTER TABLE jobs ADD COLUMN %s %s" % (col, ddl))
            ucols = {r["name"] for r in c.execute("PRAGMA table_info(usage)").fetchall()}
            if "latency" not in ucols:
                c.execute("ALTER TABLE usage ADD COLUMN latency REAL DEFAULT 0")
            jcols = {r["name"] for r in c.execute("PRAGMA table_info(jings)").fetchall()}
            for _col, _ddl in [("folder_id", "INTEGER DEFAULT 0"), ("source", "TEXT DEFAULT ''"), ("review_id", "INTEGER DEFAULT 0")]:
                if _col not in jcols:
                    c.execute("ALTER TABLE jings ADD COLUMN %s %s" % (_col, _ddl))
            icols = {r["name"] for r in c.execute("PRAGMA table_info(interviews)").fetchall()}
            for _col, _ddl in [("persona", "TEXT DEFAULT 'senior'"), ("itype", "TEXT DEFAULT '综合面'"), ("round_name", "TEXT DEFAULT '一面'"), ("avatar", "TEXT DEFAULT ''"), ("learn", "INTEGER DEFAULT 1")]:
                if _col not in icols:
                    c.execute("ALTER TABLE interviews ADD COLUMN %s %s" % (_col, _ddl))
            rcols = {r["name"] for r in c.execute("PRAGMA table_info(reviews)").fetchall()}
            if "jing_id" not in rcols:
                c.execute("ALTER TABLE reviews ADD COLUMN jing_id INTEGER DEFAULT 0")
    finally:
        c.close()


@app.on_event("startup")
def _start():
    init_db()
    key = ENV("DEEPSEEK_KEY", "")
    log.info("OfferCopilot 启动 | DEEPSEEK_KEY: %s | FEISHU_APP_ID: %s",
             (key[:6] + "...") if key else "未配置", ENV("FEISHU_APP_ID") or "未配置")
    _start_daily_loop()


PARSE_PROMPT = (
    "你是秋招岗位 JD 信息提取助手。从用户提供的 JD 文本中提取信息，"
    "严格只输出一个 JSON 对象（不要 markdown 代码块、不要任何多余文字），字段：\n"
    "company: 公司名称\n"
    "title: 岗位名称\n"
    "base: 工作地点城市，多个用、分隔\n"
    'deadline: JD 中的网申或投递截止日期，统一转成 YYYY-MM-DD；未提及则为 ""\n'
    "requirements: 岗位要求提炼为 3-6 条精简要点，每条一行，硬性要求在前\n"
)


class ParseIn(BaseModel):
    text: str


def _norm_reqs(s):
    import ast
    s = (s or "").strip()
    if s.startswith("[") and s.endswith("]"):
        try:
            arr = ast.literal_eval(s)
            if isinstance(arr, (list, tuple)):
                return "\n".join(str(x).strip() for x in arr if str(x).strip())
        except Exception:
            pass
    return s


@app.post("/api/jd/parse")
def parse_jd(p: ParseIn):
    jd = (p.text or "").strip()
    if len(jd) < 20:
        raise HTTPException(400, "JD 内容太短，请粘贴完整 JD")
    d = _llm_json(PARSE_PROMPT, jd[:6000], kind="解析")
    deadline = str(d.get("deadline", "") or "").strip()
    if not re.match(r"^\d{4}-\d{2}-\d{2}$", deadline):
        deadline = ""
    company = str(d.get("company", "") or "").strip()
    title = str(d.get("title", "") or "").strip()
    requirements = _norm_reqs(str(d.get("requirements", "") or ""))
    warn = []
    if not company:
        warn.append("未识别到公司名")
    if not title:
        warn.append("未识别到岗位名")
    if not requirements:
        warn.append("未提取到岗位要求")
    return {
        "company": company,
        "title": title,
        "base": str(d.get("base", "") or "").strip(),
        "requirements": requirements,
        "deadline": deadline,
        "warning": "；".join(warn) if warn else "",
    }


class JobIn(BaseModel):
    company: str = ""
    title: str = ""
    base: str = ""
    requirements: str = ""
    deadline: str = ""
    applied_at: str = ""
    stage: str = "未投递"
    notes: str = ""
    jd_text: str = ""


@app.get("/api/jobs")
def list_jobs():
    c = _conn()
    try:
        rows = c.execute("SELECT * FROM jobs ORDER BY id DESC").fetchall()
        return [dict(r) for r in rows]
    finally:
        c.close()


@app.post("/api/jobs")
def add_job(j: JobIn):
    if j.stage not in _all_stages():
        j.stage = "未投递"
    c = _conn()
    try:
        with c:
            cur = c.execute(
                "INSERT INTO jobs (company,title,base,requirements,deadline,applied_at,stage,notes,jd_text,created_at) VALUES (?,?,?,?,?,?,?,?,?,?)",
                (j.company, j.title, j.base, j.requirements, j.deadline, j.applied_at, j.stage, j.notes, j.jd_text, time.strftime("%Y-%m-%d %H:%M")),
            )
        return {"id": cur.lastrowid}
    finally:
        c.close()


@app.put("/api/jobs/{jid}")
def update_job(jid: int, patch: dict):
    sets, vals = [], []
    for k in JOB_FIELDS:
        if k in patch:
            v = patch[k]
            if k == "stage" and v not in _all_stages():
                raise HTTPException(400, "非法阶段: %s" % v)
            sets.append(k + "=?")
            vals.append("" if v is None else str(v))
    if not sets:
        raise HTTPException(400, "无可更新字段")
    vals.append(jid)
    c = _conn()
    try:
        with c:
            cur = c.execute("UPDATE jobs SET " + ", ".join(sets) + " WHERE id=?", vals)
        if cur.rowcount == 0:
            raise HTTPException(404, "岗位不存在")
        return {"ok": True}
    finally:
        c.close()


@app.delete("/api/jobs/{jid}")
def delete_job(jid: int):
    c = _conn()
    try:
        with c:
            cur = c.execute("DELETE FROM jobs WHERE id=?", (jid,))
        if cur.rowcount == 0:
            raise HTTPException(404, "岗位不存在")
        return {"ok": True}
    finally:
        c.close()


class ResumeIn(BaseModel):
    name: str = ""
    content: str = ""
    is_active: bool = False


def _extract_text(data: bytes, ext: str):
    ext = (ext or "").lower()
    if ext in ("md", "txt", "markdown", "text"):
        for enc in ("utf-8", "gbk"):
            try:
                return data.decode(enc)
            except Exception:
                continue
        return data.decode("utf-8", errors="ignore")
    if ext == "pdf":
        from io import BytesIO
        import pypdf
        reader = pypdf.PdfReader(BytesIO(data))
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    if ext == "docx":
        from io import BytesIO
        import docx
        document = docx.Document(BytesIO(data))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)
    raise HTTPException(400, "不支持的格式 %s（支持 pdf/docx/md/txt）" % (ext or "?"))


@app.post("/api/resume/upload")
def resume_upload(file: UploadFile = File(...)):
    name = file.filename or "简历"
    data = file.file.read()
    ext = name.rsplit(".", 1)[-1] if "." in name else ""
    try:
        text = _extract_text(data, ext)
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(400, "解析失败：%s" % e)
    if not (text or "").strip():
        raise HTTPException(400, "未能从文件中提取到文本")
    return {"name": name, "content": text.strip()}


@app.get("/api/resumes")
def list_resumes():
    c = _conn()
    try:
        rows = c.execute("SELECT * FROM resumes ORDER BY id DESC").fetchall()
        return [dict(r) for r in rows]
    finally:
        c.close()


@app.post("/api/resumes")
def add_resume(r: ResumeIn):
    c = _conn()
    try:
        with c:
            if r.is_active:
                c.execute("UPDATE resumes SET is_active=0")
            cur = c.execute(
                "INSERT INTO resumes (name,content,is_active,created_at) VALUES (?,?,?,?)",
                (r.name, r.content, 1 if r.is_active else 0, time.strftime("%Y-%m-%d %H:%M")),
            )
        return {"id": cur.lastrowid}
    finally:
        c.close()


@app.put("/api/resumes/{rid}")
def update_resume(rid: int, patch: dict):
    sets, vals = [], []
    for k in ("name", "content", "is_active"):
        if k in patch:
            v = patch[k]
            if k == "is_active":
                v = 1 if v else 0
            else:
                v = "" if v is None else str(v)
            sets.append(k + "=?")
            vals.append(v)
    if not sets:
        raise HTTPException(400, "无可更新字段")
    c = _conn()
    try:
        with c:
            row = c.execute("SELECT is_active FROM resumes WHERE id=?", (rid,)).fetchone()
            if not row:
                raise HTTPException(404, "简历不存在")
            if "is_active" in [s.split("=")[0] for s in sets] and any(v == 1 for v in vals if isinstance(v, int)):
                c.execute("UPDATE resumes SET is_active=0")
            c.execute("UPDATE resumes SET " + ", ".join(sets) + " WHERE id=?", vals + [rid])
        return {"ok": True}
    finally:
        c.close()


@app.delete("/api/resumes/{rid}")
def del_resume(rid: int):
    c = _conn()
    try:
        with c:
            cur = c.execute("DELETE FROM resumes WHERE id=?", (rid,))
        if cur.rowcount == 0:
            raise HTTPException(404, "简历不存在")
        return {"ok": True}
    finally:
        c.close()


MATCH_PROMPT = (
    "你是资深的求职匹配分析师。输入为一份简历全文和一个岗位的 JD 信息。"
    "严格只输出一个 JSON 对象（不要 markdown、不要多余文字），字段：\n"
    "overall: 综合匹配度 0-100 整数，打分严格：JD 硬性要求大部分不满足时不应超过 60\n"
    "dimensions: 对象，键为 硬性条件/技能/经历，值为 0-100 整数\n"
    "matched: 字符串数组，简历与 JD 真实匹配的亮点，只能引用简历中真实存在的内容，禁止编造\n"
    "gaps: 数组，每项 {item: 缺口描述, severity: 高/中/低}\n"
    "suggestions: 字符串数组，3-5 条，每条给出利用简历已有真实经历回应 JD 要求的具体话术，或短期可补齐的实际行动\n"
)


def _llm_cfg(override=None):
    base = _get_setting("llm_base", ENV("DEEPSEEK_BASE", "https://api.deepseek.com"))
    key = _get_setting("llm_key", ENV("DEEPSEEK_KEY", ""))
    model = _get_setting("llm_model", ENV("DEEPSEEK_MODEL", "deepseek-chat"))
    if override:
        base = override.get("base") or base
        key = override.get("key") or key
        model = override.get("model") or model
    return {"base": base.rstrip("/"), "key": key, "model": model}


def _llm_msgs(system, user, history=None):
    msgs = [{"role": "system", "content": system}]
    for h in (history or []):
        msgs.append({"role": h["role"], "content": h["content"]})
    msgs.append({"role": "user", "content": user})
    return msgs


def _llm_headers(cfg):
    headers = {
        "Authorization": "Bearer " + cfg["key"],
        "User-Agent": "offer-copilot/" + app.version,
    }
    if "opencode.ai" in cfg["base"]:
        headers["x-opencode-session"] = _LLM_SESSION
    return headers


def _cache_get(key):
    c = _conn()
    try:
        r = c.execute("SELECT value FROM cache WHERE key=?", (key,)).fetchone()
        return r["value"] if r else None
    finally:
        c.close()


def _cache_set(key, value):
    try:
        c = _conn()
        try:
            with c:
                c.execute("INSERT OR REPLACE INTO cache (key,value,created_at) VALUES (?,?,?)",
                          (key, value, time.strftime("%Y-%m-%d %H:%M")))
        finally:
            c.close()
    except Exception:
        pass


def _log_usage(kind, in_chars, out_chars, model, latency=0.0):
    try:
        c = _conn()
        try:
            with c:
                c.execute("INSERT INTO usage (kind,in_chars,out_chars,latency,model,created_at) VALUES (?,?,?,?,?,?)",
                          (kind, in_chars, out_chars, latency, model, time.strftime("%Y-%m-%d %H:%M")))
        finally:
            c.close()
    except Exception:
        pass


def _cache_key(prefix, model, system, user, history):
    payload = {"m": model, "s": system, "u": user, "h": history or []}
    return prefix + hashlib.sha256(json.dumps(payload, ensure_ascii=False, sort_keys=True).encode("utf-8")).hexdigest()


def _log_event(module, level, message):
    try:
        c = _conn()
        try:
            with c:
                c.execute("INSERT INTO op_log (module,level,message,created_at) VALUES (?,?,?,?)",
                          (module, level, str(message)[:500], time.strftime("%Y-%m-%d %H:%M:%S")))
        finally:
            c.close()
    except Exception:
        pass


def _llm_json(system, user, history=None, override=None, use_cache=True, kind="llm"):
    cfg = _llm_cfg(override)
    if not cfg["key"]:
        raise HTTPException(500, "未配置模型 API Key")
    msgs = _llm_msgs(system, user, history)
    in_chars = len(system) + len(user) + sum(len(h.get("content", "")) for h in (history or []))
    cache_key = None
    if use_cache:
        cache_key = _cache_key("json:", cfg["model"], system, user, history)
        hit = _cache_get(cache_key)
        if hit is not None:
            try:
                return json.loads(hit)
            except Exception:
                pass
    last_err = None
    t0 = time.time()
    for attempt in (1, 2):
        try:
            r = requests.post(
                cfg["base"] + "/chat/completions",
                json={
                    "model": cfg["model"],
                    "messages": msgs,
                    "response_format": {"type": "json_object"},
                    "temperature": 0.1,
                },
                headers=_llm_headers(cfg),
                timeout=240,
            )
            d = r.json()
            if r.status_code != 200 or "choices" not in d:
                err = ((d.get("error") or {}).get("message")) or d.get("message") or r.text[:200]
                raise RuntimeError("API错误: %s" % err)
            content = d["choices"][0]["message"]["content"]
            break
        except Exception as e:
            last_err = e
            if attempt == 1:
                time.sleep(2)
    else:
        log.error("模型调用失败: %s", last_err)
        _log_event("llm", "error", "模型调用失败(%s): %s" % (kind, last_err))
        raise HTTPException(502, "模型调用失败：%s" % last_err)
    _log_usage(kind, in_chars, len(content), cfg["model"], round(time.time() - t0, 2))
    try:
        m = re.search(r"\{.*\}", content, re.S)
        result = json.loads(m.group(0) if m else content)
    except Exception:
        raise HTTPException(502, "模型输出异常，请重试")
    if cache_key:
        _cache_set(cache_key, json.dumps(result, ensure_ascii=False))
    return result


def _llm_text(system, user, override=None, use_cache=True, kind="llm"):
    cfg = _llm_cfg(override)
    if not cfg["key"]:
        raise HTTPException(500, "未配置模型 API Key")
    in_chars = len(system) + len(user)
    cache_key = None
    if use_cache:
        cache_key = _cache_key("txt:", cfg["model"], system, user, None)
        hit = _cache_get(cache_key)
        if hit is not None:
            return hit
    last_err = None
    t0 = time.time()
    for attempt in (1, 2):
        try:
            r = requests.post(
                cfg["base"] + "/chat/completions",
                json={
                    "model": cfg["model"],
                    "messages": _llm_msgs(system, user),
                    "temperature": 0.4,
                },
                headers=_llm_headers(cfg),
                timeout=240,
            )
            d = r.json()
            if r.status_code != 200 or "choices" not in d:
                err = ((d.get("error") or {}).get("message")) or d.get("message") or r.text[:200]
                raise RuntimeError("API错误: %s" % err)
            content = d["choices"][0]["message"]["content"]
            break
        except Exception as e:
            last_err = e
            if attempt == 1:
                time.sleep(2)
    else:
        log.error("模型调用失败: %s", last_err)
        _log_event("llm", "error", "模型调用失败(%s): %s" % (kind, last_err))
        raise HTTPException(502, "模型调用失败：%s" % last_err)
    _log_usage(kind, in_chars, len(content), cfg["model"], round(time.time() - t0, 2))
    if cache_key:
        _cache_set(cache_key, content)
    return content


def _llm_stream(system, user, history=None, override=None, kind="llm"):
    cfg = _llm_cfg(override)
    if not cfg["key"]:
        raise HTTPException(500, "未配置模型 API Key")
    in_chars = len(system) + len(user) + sum(len(h.get("content", "")) for h in (history or []))
    t0 = time.time()
    r = requests.post(
        cfg["base"] + "/chat/completions",
        json={
            "model": cfg["model"],
            "messages": _llm_msgs(system, user, history),
            "stream": True,
            "temperature": 0.4,
        },
        headers=_llm_headers(cfg),
        timeout=180,
        stream=True,
    )
    if r.status_code != 200:
        _log_event("llm", "error", "模型调用失败(%s): %s" % (kind, r.text[:200]))
        raise HTTPException(502, "模型调用失败: %s" % r.text[:200])
    out = []
    for line in r.iter_lines():
        if not line:
            continue
        line = line.decode("utf-8", errors="ignore").strip()
        if not line.startswith("data:"):
            continue
        data = line[5:].strip()
        if data == "[DONE]":
            break
        try:
            chunk = json.loads(data)
            delta = chunk["choices"][0].get("delta", {}).get("content")
            if delta:
                out.append(delta)
                yield delta
        except Exception:
            continue
    _log_usage(kind, in_chars, sum(len(x) for x in out), cfg["model"], round(time.time() - t0, 2))


def _clamp(v, lo=0, hi=100):
    try:
        v = int(v)
    except Exception:
        v = lo
    return max(lo, min(hi, v))


@app.post("/api/jobs/{jid}/match")
def match_job(jid: int, body: dict = None):
    body = body or {}
    c = _conn()
    try:
        job = c.execute("SELECT * FROM jobs WHERE id=?", (jid,)).fetchone()
    finally:
        c.close()
    if not job:
        raise HTTPException(404, "岗位不存在")
    rid = body.get("resume_id")
    c = _conn()
    try:
        if rid:
            resume = c.execute("SELECT * FROM resumes WHERE id=?", (rid,)).fetchone()
        else:
            resume = c.execute("SELECT * FROM resumes WHERE is_active=1 ORDER BY id DESC").fetchone()
    finally:
        c.close()
    if not resume:
        raise HTTPException(400, "请先在「简历」中添加并启用一份简历")
    if not (job["requirements"] or "").strip() and not (job["jd_text"] or "").strip():
        raise HTTPException(400, "该岗位缺少要求与 JD 信息，无法匹配")
    res = _run_evaluation(job, resume)
    evaluation = res["evaluation"]
    overall = _clamp(evaluation.get("final_score", evaluation.get("raw_score", 0)))
    c = _conn()
    try:
        with c:
            c.execute(
                "UPDATE jobs SET match_score=?, match_detail=?, matched_at=? WHERE id=?",
                (overall, json.dumps(evaluation, ensure_ascii=False), time.strftime("%Y-%m-%d %H:%M"), jid),
            )
    finally:
        c.close()
    return {"ok": True, "overall": overall, "detail": evaluation}


PERSONAS = {
    "peer": {"name": "小林", "style": "一线工程师面试官，随和、朋友式，侧重实操细节、具体做法和踩坑经历"},
    "senior": {"name": "陈工", "style": "资深技术面试官，严谨、逻辑性强，喜欢深挖原理、架构和系统设计，追问尖锐"},
    "hr": {"name": "王HR", "style": "HR 面试官，关注求职动机、岗位匹配度、软技能、稳定性和薪资期望"},
    "manager": {"name": "刘总", "style": "部门经理，沉稳克制，考察全局观、业务思维、跨团队协作和成长潜力"},
}

ITYPE_GUIDE = {
    "技术面": "技术面试，侧重技术原理、项目深挖、系统设计、难点取舍",
    "行为面": "行为面试，侧重 STAR 行为题、软技能、团队协作、情景应变",
    "HR面": "HR 面试，侧重求职动机、岗位匹配、稳定性、薪资与职业规划",
    "综合面": "综合面试，技术、行为、动机均衡覆盖",
}


def _iv_system(job, resume_text, persona, itype, round_name, jing_qs, style_profile=""):
    p = PERSONAS.get(persona or "senior", PERSONAS["senior"])
    guide = ITYPE_GUIDE.get(itype or "综合面", ITYPE_GUIDE["综合面"])
    rn = round_name or "一面"
    prompt = (
        "你是%s，%s。\n"
        "正在对候选人进行【%s】的%s模拟面试。\n"
        "规则：\n"
        "1. 首轮输出：简短自我介绍开场（你是谁、这场面侧重什么）+ 第一个问题（请候选人做个自我介绍）\n"
        "2. 之后每轮输出 = 对候选人上一条回答的简短点评(reply) + 本题评分(score 0-100整数，打分严格) + 一句具体反馈(feedback) + 下一个问题(question)\n"
        "3. 追问要引用候选人回答的原话，深挖 STAR 细节（具体做了什么、量化结果、困难与取舍）\n"
        "4. 计划 5 个左右主问题，每个主问题最多追问 1 次\n"
        "5. 全部问完后 kind=end，question 写「我的问题问完了，可以点击结束生成报告」\n"
        "6. 用中文，语气符合你的人设，一次只问一个问题\n"
        "严格只输出 JSON：{reply: string, score: int或null, feedback: string, question: string, kind: followup/new/end}\n"
    ) % (p["name"], p["style"], rn, guide)
    prompt += "\n\n" + _job_ctx(job)
    prompt += "\n\n【候选人简历】\n" + (resume_text[:6000] or "（未提供简历）")
    if jing_qs:
        prompt += "\n\n【该公司真实/相关面试问题（优先参考，尽量据此出题/追问）】\n" + "\n".join(jing_qs)
    if style_profile:
        prompt += "\n\n【候选人真实面试风格画像（模拟面试应尽量贴近这种风格）】\n" + style_profile
    return prompt


REPORT_PROMPT = (
    "你是资深面试评估官，基于岗位信息、候选人简历和完整面试对话，生成客观、准确、公正的候选人评估报告。\n"
    "评分与评价必须严格遵守以下规则：\n"
    "1. 只依据面试对话转写中真实存在的内容评价；每条反馈/结论都要能在转写中找到依据，必要时引用转写原文\n"
    "2. 禁止臆断：不得声称候选人\"没提到/未涉及/没有说\"某内容，除非转写中确实完全没有该内容；不得编造候选人说过的话\n"
    "3. 简历对比：仅当简历已提供时对照简历；只指出简历与转写确实矛盾的地方；候选人在转写中已说明的实现方式、细节与简历一致即视为一致，不得误判为\"与简历不符\"\n"
    "4. 打分公正：基于转写中的实际表现证据综合评分；对表述清楚但可进一步展开的内容给予正常分数而非判缺失；个别小疏漏不应大幅拉低总分；overall 反映整体水平而非单点\n"
    "5. 岗位匹配以 JD 要求为参照，缺依据时不强行扣分\n"
    "严格只输出 JSON（不要 markdown、不要多余文字）：\n"
    "overall: 综合得分 0-100 整数\n"
    "dimensions: 对象，键为 专业深度/表达与结构/岗位匹配/思考深度，值 0-100 整数\n"
    "per_q: 数组，每项 {question: 问题摘要, verdict: 好/中/差, feedback: 1-2 句具体反馈（须引用转写依据）}\n"
    "strengths: 字符串数组 2-4 条亮点（须有转写依据）\n"
    "improvements: 字符串数组 2-4 条待改进（须有转写依据）\n"
    "practice: 字符串数组 1-3 条接下来的练习建议\n"
)


def _get_job(jid):
    c = _conn()
    try:
        return c.execute("SELECT * FROM jobs WHERE id=?", (jid,)).fetchone()
    finally:
        c.close()


def _job_ctx(job):
    return "【岗位信息】\n公司：%s\n岗位：%s\nBase：%s\n要求：\n%s\n\n【JD 原文】\n%s" % (
        job["company"], job["title"], job["base"], job["requirements"] or "（无）", (job["jd_text"] or "（无）")[:3000])


@app.get("/api/interviews")
def list_interviews():
    c = _conn()
    try:
        rows = c.execute(
            "SELECT id, job_id, title, status, overall, created_at, finished_at FROM interviews ORDER BY id DESC"
        ).fetchall()
        return [dict(r) for r in rows]
    finally:
        c.close()


@app.get("/api/interviews/{iid}")
def get_interview(iid: int):
    c = _conn()
    try:
        row = c.execute("SELECT * FROM interviews WHERE id=?", (iid,)).fetchone()
    finally:
        c.close()
    if not row:
        raise HTTPException(404, "面试不存在")
    d = dict(row)
    d["messages"] = json.loads(d.get("messages") or "[]")
    d["report"] = json.loads(d.get("report") or "null")
    d["persona_name"] = PERSONAS.get(d.get("persona") or "senior", PERSONAS["senior"])["name"]
    return d


@app.delete("/api/interviews/{iid}")
def del_interview(iid: int):
    c = _conn()
    try:
        with c:
            cur = c.execute("DELETE FROM interviews WHERE id=?", (iid,))
        if cur.rowcount == 0:
            raise HTTPException(404, "面试不存在")
        return {"ok": True}
    finally:
        c.close()


@app.post("/api/interview/start")
def interview_start(body: dict):
    jid = body.get("job_id")
    if isinstance(jid, str):
        jid = int(jid) if jid.isdigit() else None
    job = _get_job(jid) if jid else None
    if not job:
        raise HTTPException(404, "岗位不存在")
    persona = str(body.get("persona") or "senior").strip()
    itype = str(body.get("itype") or "综合面").strip()
    round_name = str(body.get("round") or "一面").strip()
    avatar = str(body.get("avatar") or "").strip()[:200]
    learn = 1 if body.get("learn", True) else 0
    c = _conn()
    try:
        resume = c.execute("SELECT content FROM resumes WHERE is_active=1 ORDER BY id DESC").fetchone()
    finally:
        c.close()
    resume_text = (resume["content"] if resume else "") or ""
    title = (job["company"] or "?") + " · " + (job["title"] or "?")
    jing_qs = _interview_q_context(job["company"], bool(learn))
    style_profile = _interview_style_profile() if learn else ""
    prompt = _iv_system(job, resume_text, persona, itype, round_name, jing_qs, style_profile)
    d = _llm_json(prompt, "面试开始，请按规则输出首轮（自我介绍开场 + 第一个问题）。", override=_eval_cfg(), kind="面试")
    try:
        score = int(d.get("score")) if d.get("score") is not None else None
    except Exception:
        score = None
    first = {
        "role": "assistant",
        "reply": str(d.get("reply", "") or "").strip(),
        "score": score,
        "feedback": str(d.get("feedback", "") or "").strip(),
        "question": str(d.get("question", "") or "").strip() or "请先用 1 分钟做个自我介绍。",
        "kind": "new",
    }
    c = _conn()
    try:
        with c:
            cur = c.execute(
                "INSERT INTO interviews (job_id,title,resume_text,messages,report,status,overall,persona,itype,round_name,avatar,learn,created_at) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)",
                (jid, title, resume_text, json.dumps([first], ensure_ascii=False), "", "ing", 0, persona, itype, round_name, avatar, learn, time.strftime("%Y-%m-%d %H:%M")),
            )
        return {"id": cur.lastrowid, "title": title, "messages": [first], "avatar": avatar,
                "persona_name": PERSONAS.get(persona, PERSONAS["senior"])["name"]}
    finally:
        c.close()


@app.post("/api/interview/{iid}/answer")
def interview_answer(iid: int, body: dict):
    content = (body.get("content") or "").strip()
    if not content:
        raise HTTPException(400, "回答不能为空")
    c = _conn()
    try:
        row = c.execute("SELECT * FROM interviews WHERE id=?", (iid,)).fetchone()
    finally:
        c.close()
    if not row:
        raise HTTPException(404, "面试不存在")
    if row["status"] != "ing":
        raise HTTPException(400, "面试已结束，不能继续作答")
    msgs = json.loads(row["messages"] or "[]")
    msgs.append({"role": "user", "content": content[:4000]})
    job = _get_job(row["job_id"])
    jing_qs = _interview_q_context(job["company"], bool(row["learn"]))
    style_profile = _interview_style_profile() if row["learn"] else ""
    prompt = _iv_system(job, row["resume_text"], row["persona"], row["itype"], row["round_name"], jing_qs, style_profile)
    history = []
    for m in msgs[:-1]:
        if m.get("role") == "user":
            history.append({"role": "user", "content": m.get("content", "")})
        else:
            history.append({"role": "assistant", "content": ((m.get("reply", "") or "") + "\n" + (m.get("question", "") or "")).strip()})
    d = _llm_json(prompt, "候选人最新回答：\n" + content + "\n\n请输出 JSON（reply + score + feedback + question + kind）。", history, override=_eval_cfg(), kind="面试")
    kind = str(d.get("kind", "") or "").strip()
    if kind not in ("followup", "new", "end"):
        kind = "new"
    try:
        score = int(d.get("score")) if d.get("score") is not None else None
        if score is not None:
            score = max(0, min(100, score))
    except Exception:
        score = None
    new = {
        "role": "assistant",
        "reply": str(d.get("reply", "") or "").strip(),
        "score": score,
        "feedback": str(d.get("feedback", "") or "").strip(),
        "question": str(d.get("question", "") or "").strip() or "请继续说说你的项目。",
        "kind": kind,
    }
    msgs.append(new)
    c = _conn()
    try:
        with c:
            c.execute("UPDATE interviews SET messages=? WHERE id=?", (json.dumps(msgs, ensure_ascii=False), iid))
    finally:
        c.close()
    return {"message": new}


@app.post("/api/interview/{iid}/finish")
def interview_finish(iid: int):
    c = _conn()
    try:
        row = c.execute("SELECT * FROM interviews WHERE id=?", (iid,)).fetchone()
    finally:
        c.close()
    if not row:
        raise HTTPException(404, "面试不存在")
    if row["status"] == "done" and row["report"]:
        return {"report": json.loads(row["report"])}
    msgs = json.loads(row["messages"] or "[]")
    if not [m for m in msgs if m.get("role") == "user"]:
        raise HTTPException(400, "还没有作答记录，无法生成报告")
    job = _get_job(row["job_id"])
    lines = []
    for m in msgs:
        if m.get("role") == "user":
            lines.append("候选人：" + m.get("content", ""))
        else:
            line = ((m.get("reply", "") or "") + " " + (m.get("question", "") or "")).strip()
            if m.get("score") is not None:
                line += "（本题评分：%s；反馈：%s）" % (m["score"], m.get("feedback", "") or "")
            lines.append("面试官：" + line)
    transcript = "\n".join(lines)[:8000]
    user = (_job_ctx(job) + "\n\n【候选人简历】\n" + (row["resume_text"] or "（未提供简历）") +
            "\n\n【面试对话转写】\n" + transcript)
    d = _llm_json(REPORT_PROMPT, user, override=_eval_cfg(), kind="面试报告")
    overall = _clamp(d.get("overall", 0))
    dims = {}
    for k, v in (d.get("dimensions") or {}).items():
        dims[str(k)] = _clamp(v)
    per_q = []
    for q in (d.get("per_q") or []):
        if isinstance(q, dict):
            verdict = str(q.get("verdict", "中")).strip()
            per_q.append({
                "question": str(q.get("question", "")).strip(),
                "verdict": verdict if verdict in ("好", "中", "差") else "中",
                "feedback": str(q.get("feedback", "")).strip(),
            })
    report = {
        "overall": overall,
        "dimensions": dims,
        "per_q": per_q[:12],
        "strengths": [str(x).strip() for x in (d.get("strengths") or []) if str(x).strip()][:6],
        "improvements": [str(x).strip() for x in (d.get("improvements") or []) if str(x).strip()][:6],
        "practice": [str(x).strip() for x in (d.get("practice") or []) if str(x).strip()][:4],
    }
    c = _conn()
    try:
        with c:
            c.execute(
                "UPDATE interviews SET report=?, status='done', overall=?, finished_at=? WHERE id=?",
                (json.dumps(report, ensure_ascii=False), overall, time.strftime("%Y-%m-%d %H:%M"), iid),
            )
    finally:
        c.close()
    if row["learn"]:
        _save_mock_jing(job, msgs)
    return {"report": report}


def _norm(s):
    return re.sub(r"\s+", "", s or "")


@app.get("/api/evalsets")
def list_evalsets():
    c = _conn()
    try:
        rows = c.execute("SELECT * FROM evalsets ORDER BY id DESC").fetchall()
        out = []
        for r in rows:
            n = c.execute("SELECT COUNT(*) FROM evalcases WHERE evalset_id=?", (r["id"],)).fetchone()[0]
            d = dict(r)
            d["case_count"] = n
            out.append(d)
        return out
    finally:
        c.close()


class EvalSetIn(BaseModel):
    name: str = ""


@app.post("/api/evalsets")
def add_evalset(e: EvalSetIn):
    c = _conn()
    try:
        with c:
            cur = c.execute(
                "INSERT INTO evalsets (name,created_at) VALUES (?,?)",
                (e.name or "未命名评测集", time.strftime("%Y-%m-%d %H:%M")),
            )
        return {"id": cur.lastrowid}
    finally:
        c.close()


@app.delete("/api/evalsets/{sid}")
def del_evalset(sid: int):
    c = _conn()
    try:
        with c:
            cur = c.execute("DELETE FROM evalsets WHERE id=?", (sid,))
            c.execute("DELETE FROM evalcases WHERE evalset_id=?", (sid,))
            c.execute("DELETE FROM eval_runs WHERE evalset_id=?", (sid,))
        if cur.rowcount == 0:
            raise HTTPException(404, "评测集不存在")
        return {"ok": True}
    finally:
        c.close()


@app.post("/api/evalsets/{sid}/cases_from_jobs")
def evalset_from_jobs(sid: int):
    c = _conn()
    try:
        if not c.execute("SELECT id FROM evalsets WHERE id=?", (sid,)).fetchone():
            raise HTTPException(404, "评测集不存在")
        jobs = c.execute("SELECT * FROM jobs WHERE jd_text IS NOT NULL AND length(jd_text) > 20").fetchall()
        added = 0
        existing = {r["jd_text"] for r in c.execute("SELECT jd_text FROM evalcases WHERE evalset_id=?", (sid,)).fetchall()}
        with c:
            for j in jobs:
                if j["jd_text"] in existing:
                    continue
                c.execute(
                    "INSERT INTO evalcases (evalset_id,jd_text,company,title,base,deadline,created_at) VALUES (?,?,?,?,?,?,?)",
                    (sid, j["jd_text"], j["company"], j["title"], j["base"], j["deadline"], time.strftime("%Y-%m-%d %H:%M")),
                )
                added += 1
        return {"added": added}
    finally:
        c.close()


@app.get("/api/evalsets/{sid}/cases")
def list_cases(sid: int):
    c = _conn()
    try:
        rows = c.execute("SELECT * FROM evalcases WHERE evalset_id=? ORDER BY id DESC", (sid,)).fetchall()
        return [dict(r) for r in rows]
    finally:
        c.close()


class EvalCaseIn(BaseModel):
    evalset_id: int = 0
    jd_text: str = ""
    company: str = ""
    title: str = ""
    base: str = ""
    deadline: str = ""


@app.post("/api/evalcases")
def add_case(e: EvalCaseIn):
    if not (e.jd_text or "").strip():
        raise HTTPException(400, "JD 不能为空")
    c = _conn()
    try:
        with c:
            cur = c.execute(
                "INSERT INTO evalcases (evalset_id,jd_text,company,title,base,deadline,created_at) VALUES (?,?,?,?,?,?,?)",
                (e.evalset_id, e.jd_text, e.company, e.title, e.base, e.deadline, time.strftime("%Y-%m-%d %H:%M")),
            )
        return {"id": cur.lastrowid}
    finally:
        c.close()


@app.put("/api/evalcases/{cid}")
def update_case(cid: int, patch: dict):
    fields = ("jd_text", "company", "title", "base", "deadline")
    sets, vals = [], []
    for k in fields:
        if k in patch:
            sets.append(k + "=?")
            vals.append("" if patch[k] is None else str(patch[k]))
    if not sets:
        raise HTTPException(400, "无可更新字段")
    vals.append(cid)
    c = _conn()
    try:
        with c:
            cur = c.execute("UPDATE evalcases SET " + ", ".join(sets) + " WHERE id=?", vals)
        if cur.rowcount == 0:
            raise HTTPException(404, "用例不存在")
        return {"ok": True}
    finally:
        c.close()


@app.delete("/api/evalcases/{cid}")
def del_case(cid: int):
    c = _conn()
    try:
        with c:
            cur = c.execute("DELETE FROM evalcases WHERE id=?", (cid,))
        if cur.rowcount == 0:
            raise HTTPException(404, "用例不存在")
        return {"ok": True}
    finally:
        c.close()


@app.get("/api/prompt_versions")
def list_prompt_versions():
    c = _conn()
    try:
        rows = c.execute("SELECT * FROM prompt_versions ORDER BY id ASC").fetchall()
        return [dict(r) for r in rows]
    finally:
        c.close()


class PromptVerIn(BaseModel):
    name: str = ""
    prompt: str = ""


@app.post("/api/prompt_versions")
def add_prompt_version(p: PromptVerIn):
    if not (p.prompt or "").strip():
        raise HTTPException(400, "Prompt 不能为空")
    c = _conn()
    try:
        with c:
            cur = c.execute(
                "INSERT INTO prompt_versions (name,prompt,created_at) VALUES (?,?,?)",
                (p.name or "未命名版本", p.prompt, time.strftime("%Y-%m-%d %H:%M")),
            )
        return {"id": cur.lastrowid}
    finally:
        c.close()


@app.delete("/api/prompt_versions/{vid}")
def del_prompt_version(vid: int):
    c = _conn()
    try:
        row = c.execute("SELECT name FROM prompt_versions WHERE id=?", (vid,)).fetchone()
        if not row:
            raise HTTPException(404, "版本不存在")
        if row["name"] == DEFAULT_PROMPT_NAME:
            raise HTTPException(400, "默认模板不可删除")
        with c:
            c.execute("DELETE FROM prompt_versions WHERE id=?", (vid,))
        return {"ok": True}
    finally:
        c.close()


EVAL_FIELDS = ("company", "title", "base", "deadline")


def _run_eval(sid, vid, model_ov, consistency):
    c = _conn()
    try:
        cases = c.execute("SELECT * FROM evalcases WHERE evalset_id=? ORDER BY id DESC", (sid,)).fetchall()
        if not cases:
            raise HTTPException(400, "评测集为空，先添加或用例导入岗位")
        if vid:
            version = c.execute("SELECT * FROM prompt_versions WHERE id=?", (vid,)).fetchone()
            prompt = version["prompt"] if version else PARSE_PROMPT
        else:
            prompt = PARSE_PROMPT
    finally:
        c.close()
    counts = {f: 0 for f in EVAL_FIELDS}
    totals = {f: 0 for f in EVAL_FIELDS}
    stables = {f: 0 for f in EVAL_FIELDS}
    recall_num = {f: 0 for f in EVAL_FIELDS}
    recall_den = {f: 0 for f in EVAL_FIELDS}
    per_case = []
    for case in cases:
        parses = []
        for _ in range(consistency):
            d = _llm_json(prompt, (case["jd_text"] or "")[:6000], override=model_ov, use_cache=False, kind="评测")
            p = {f: str(d.get(f, "") or "").strip() for f in EVAL_FIELDS}
            if not re.match(r"^\d{4}-\d{2}-\d{2}$", p["deadline"]):
                p["deadline"] = ""
            parses.append(p)
        res = {}
        for f in EVAL_FIELDS:
            has_golden = (case[f] or "").strip() != ""
            correct = sum(1 for p in parses if _norm(p[f]) == _norm(case[f]))
            ok = correct >= (consistency // 2 + 1)
            if f == "deadline" and not has_golden:
                res[f] = ok
            else:
                totals[f] += 1
                if ok:
                    counts[f] += 1
                res[f] = ok
            st = len({_norm(p[f]) for p in parses}) == 1
            res[f + "_stable"] = st
            if st:
                stables[f] += 1
            if has_golden:
                recall_den[f] += 1
                if ok:
                    recall_num[f] += 1
        per_case.append({
            "id": case["id"],
            "jd": (case["jd_text"] or "")[:36],
            "golden": {"company": case["company"], "title": case["title"], "base": case["base"], "deadline": case["deadline"]},
            "parsed": parses[0],
            "res": res,
            "n": consistency,
        })
    summary = {}
    for f in EVAL_FIELDS:
        summary[f] = round(100.0 * counts[f] / totals[f], 1) if totals[f] else None
        summary[f + "_stable"] = round(100.0 * stables[f] / totals[f], 1) if totals[f] else None
        summary[f + "_recall"] = round(100.0 * recall_num[f] / recall_den[f], 1) if recall_den[f] else None
    avgs = [v for v in [summary[f] for f in EVAL_FIELDS] if v is not None]
    summary["avg"] = round(sum(avgs) / len(avgs), 1) if avgs else None
    savgs = [v for v in [summary[f + "_stable"] for f in EVAL_FIELDS] if v is not None]
    summary["stable_avg"] = round(sum(savgs) / len(savgs), 1) if savgs else None
    ravs = [v for v in [summary[f + "_recall"] for f in EVAL_FIELDS] if v is not None]
    summary["recall_avg"] = round(sum(ravs) / len(ravs), 1) if ravs else None
    summary["model"] = (model_ov or {}).get("model") or _llm_cfg()["model"]
    summary["consistency"] = consistency
    return {"summary": summary, "per_case": per_case}


@app.post("/api/eval/run")
def eval_run(body: dict):
    sid = body.get("evalset_id")
    vid = body.get("version_id")
    consistency = int(body.get("consistency") or 1)
    consistency = max(1, min(5, consistency))
    model_ov = body.get("model") or None
    r = _run_eval(sid, vid, model_ov, consistency)
    c = _conn()
    try:
        with c:
            cur = c.execute(
                "INSERT INTO eval_runs (evalset_id,version_id,per_case,summary,created_at) VALUES (?,?,?,?,?)",
                (sid, vid, json.dumps(r["per_case"], ensure_ascii=False), json.dumps(r["summary"], ensure_ascii=False), time.strftime("%Y-%m-%d %H:%M")),
            )
        return {"id": cur.lastrowid, "summary": r["summary"], "per_case": r["per_case"]}
    finally:
        c.close()


@app.post("/api/eval/compare")
def eval_compare(body: dict):
    sid = body.get("evalset_id")
    vid = body.get("version_id")
    models = [{"name": "当前模型", "cfg": _llm_cfg()}]
    try:
        saved = json.loads(_get_setting("saved_models", "[]") or "[]")
    except Exception:
        saved = []
    for s in saved:
        if isinstance(s, dict) and (s.get("base") or s.get("model")):
            models.append({
                "name": s.get("name") or s.get("model") or "?",
                "cfg": {"base": s.get("base", ""), "key": s.get("key", ""), "model": s.get("model", "")},
            })
    results = []
    for m in models:
        try:
            r = _run_eval(sid, vid, m["cfg"], 1)
            results.append({"name": m["name"], "summary": r["summary"], "error": None})
        except Exception as e:
            results.append({"name": m["name"], "summary": None, "error": str(e)[:150]})
    return {"results": results}


@app.get("/api/eval_runs")
def list_eval_runs():
    c = _conn()
    try:
        rows = c.execute("SELECT * FROM eval_runs ORDER BY id DESC LIMIT 50").fetchall()
        out = []
        for r in rows:
            d = dict(r)
            d["summary"] = json.loads(r["summary"] or "null")
            name = c.execute("SELECT name FROM evalsets WHERE id=?", (r["evalset_id"],)).fetchone()
            vname = c.execute("SELECT name FROM prompt_versions WHERE id=?", (r["version_id"],)).fetchone() if r["version_id"] else None
            d["set_name"] = name["name"] if name else "?"
            d["ver_name"] = vname["name"] if vname else DEFAULT_PROMPT_NAME
            out.append(d)
        return out
    finally:
        c.close()


@app.get("/api/eval_runs/{rid}")
def get_eval_run(rid: int):
    c = _conn()
    try:
        row = c.execute("SELECT * FROM eval_runs WHERE id=?", (rid,)).fetchone()
    finally:
        c.close()
    if not row:
        raise HTTPException(404, "运行不存在")
    d = dict(row)
    d["summary"] = json.loads(row["summary"] or "null")
    d["per_case"] = json.loads(row["per_case"] or "[]")
    return d


@app.delete("/api/eval_runs/{rid}")
def del_eval_run(rid: int):
    c = _conn()
    try:
        with c:
            cur = c.execute("DELETE FROM eval_runs WHERE id=?", (rid,))
        if cur.rowcount == 0:
            raise HTTPException(404, "运行不存在")
        return {"ok": True}
    finally:
        c.close()


def _get_setting(key, default=""):
    c = _conn()
    try:
        r = c.execute("SELECT value FROM settings WHERE key=?", (key,)).fetchone()
        return r["value"] if r else default
    finally:
        c.close()


def _set_setting(key, value):
    c = _conn()
    try:
        with c:
            c.execute("INSERT OR REPLACE INTO settings (key,value) VALUES (?,?)", (key, str(value)))
    finally:
        c.close()


def _days_left(d):
    if not d:
        return None
    try:
        due = time.mktime(time.strptime(d, "%Y-%m-%d"))
        return int((due - time.time()) // 86400)
    except Exception:
        return None


def _fmt_job(j):
    s = "%s · %s" % (j["company"] or "?", j["title"] or "?")
    if j["base"]:
        s += "（%s）" % j["base"]
    if j["deadline"]:
        s += " 截止 %s" % j["deadline"]
    if j["match_score"]:
        s += " 匹配度 %d" % j["match_score"]
    return s


def _build_digest_data():
    c = _conn()
    try:
        jobs = [dict(r) for r in c.execute("SELECT * FROM jobs ORDER BY id DESC").fetchall()]
    finally:
        c.close()
    today = time.strftime("%Y-%m-%d")
    new_list, urgent_list, high_list = [], [], []
    stage_count = {}
    for j in jobs:
        stage = j["stage"] or "未投递"
        stage_count[stage] = stage_count.get(stage, 0) + 1
        if (j["created_at"] or "")[:10] == today:
            new_list.append(_fmt_job(j))
        if stage not in ("Offer", "已挂"):
            dl = _days_left(j["deadline"])
            if dl is not None and 0 <= dl <= 3:
                urgent_list.append(_fmt_job(j) + "（%d 天后截止）" % dl)
            if (j["match_score"] or 0) >= 70:
                high_list.append(_fmt_job(j))
    parts = [
        "【今日新增岗位】\n" + ("\n".join(new_list) if new_list else "无"),
        "【即将截止（3 天内）】\n" + ("\n".join(urgent_list) if urgent_list else "无"),
        "【高匹配机会（匹配度≥70）】\n" + ("\n".join(high_list) if high_list else "无"),
    ]
    order = ["未投递", "已投递", "简历评估", "AI面试", "一面", "二面", "三面", "HR面", "Offer", "已挂"]
    sc = "、".join("%s %d" % (k, stage_count[k]) for k in order if stage_count.get(k))
    parts.append("【进度总览】\n" + (sc or "暂无"))
    return "\n\n".join(parts)


DIGEST_PROMPT = (
    "你是求职日报编辑。基于下面提供的岗位数据，生成一份简洁有信息量的《秋招日报》。"
    "用中文，可用少量 emoji 作为分区标题，总字数控制在 500 字内。"
    "保留数据里的关键信息（公司、岗位、地点、截止时间、匹配度），不要编造新岗位。"
    "没有数据的板块直接省略。结尾附一句行动建议。直接输出日报正文，不要多余解释。"
)


@app.post("/api/digest/generate")
def digest_generate():
    data = _build_digest_data()
    text = _llm_text(DIGEST_PROMPT, data, kind="日报").strip()
    return {"text": text}


@app.post("/api/digest/push")
def digest_push(body: dict = None):
    body = body or {}
    text = (body.get("text") or "").strip()
    if not text:
        text = digest_generate()["text"]
    chat_id = body.get("chat_id") or _get_setting("feishu_chat_id", ENV("FEISHU_CHAT_ID", ""))
    if not chat_id:
        raise HTTPException(400, "未配置飞书 chat_id，请在「日报」页填写并保存")
    if not ENV("FEISHU_APP_ID") or not ENV("FEISHU_APP_SECRET"):
        raise HTTPException(400, "未配置飞书应用凭证（FEISHU_APP_ID/SECRET）")
    try:
        feishu.send_text(chat_id, text)
        _log_event("feishu", "info", "日报推送成功 -> %s" % chat_id)
    except Exception as e:
        _log_event("feishu", "error", "日报推送失败: %s" % e)
        raise HTTPException(502, "推送失败：%s" % e)
    return {"ok": True}


@app.get("/api/digest/config")
def digest_config():
    return {
        "chat_id": _get_setting("feishu_chat_id", ENV("FEISHU_CHAT_ID", "")),
        "enabled": _get_setting("daily_enabled", ENV("FEISHU_DAILY_ENABLED", "0")) == "1",
        "time": _get_setting("daily_time", ENV("FEISHU_DAILY_TIME", "09:00")),
        "has_credentials": bool(ENV("FEISHU_APP_ID") and ENV("FEISHU_APP_SECRET")),
    }


@app.post("/api/digest/config")
def digest_set_config(body: dict):
    for k, v in body.items():
        if k == "chat_id":
            _set_setting("feishu_chat_id", v)
        elif k == "enabled":
            _set_setting("daily_enabled", "1" if v else "0")
        elif k == "time":
            _set_setting("daily_time", v)
    return {"ok": True}


_daily_done = set()


def _daily_loop():
    while True:
        try:
            if _get_setting("daily_enabled", ENV("FEISHU_DAILY_ENABLED", "0")) == "1":
                now = time.strftime("%H:%M")
                target = _get_setting("daily_time", ENV("FEISHU_DAILY_TIME", "09:00"))
                today = time.strftime("%Y-%m-%d")
                if now == target and today not in _daily_done:
                    chat_id = _get_setting("feishu_chat_id", ENV("FEISHU_CHAT_ID", ""))
                    if chat_id:
                        text = digest_generate()["text"]
                        feishu.send_text(chat_id, text)
                        _daily_done.add(today)
                        log.info("日报已定时推送 %s", today)
        except Exception as e:
            log.warning("日报定时任务异常: %s", e)
        time.sleep(60)


def _start_daily_loop():
    threading.Thread(target=_daily_loop, daemon=True).start()


def _active_resume():
    c = _conn()
    try:
        return c.execute("SELECT * FROM resumes WHERE is_active=1 ORDER BY id DESC").fetchone()
    finally:
        c.close()


def _job_and_resume(body):
    jid = body.get("job_id")
    if isinstance(jid, str):
        jid = int(jid) if jid.isdigit() else None
    job = _get_job(jid) if jid else None
    if not job:
        raise HTTPException(404, "岗位不存在")
    resume = _active_resume()
    if not resume:
        raise HTTPException(400, "请先在「简历」中添加并启用一份简历")
    return job, resume


def _mat_user_content(job, resume):
    return ("【岗位信息】\n公司：%s\n岗位：%s\nBase：%s\n要求：\n%s\n\n【JD 原文】\n%s\n\n【简历全文】\n%s") % (
        job["company"], job["title"], job["base"], job["requirements"] or "（无）",
        (job["jd_text"] or "（无）")[:3000], (resume["content"] or "")[:6000],
    )


TAILOR_PROMPT = (
    "你是资深简历顾问。基于候选人简历和岗位 JD，生成一份针对该岗位定制的简历。\n"
    "规则：\n"
    "1. 只能改写、重组、强调简历中真实存在的内容，禁止编造经历、技能、数据、奖项\n"
    "2. 用 STAR 法则改写项目与经历 bullet，把 JD 关键词自然注入真实经历\n"
    "3. 结构与原简历保持一致，输出一份完整、可直接投递的 Markdown 简历\n"
    "严格只输出 JSON：{resume_md: string(完整Markdown简历), changes: [string](针对该岗位做的关键调整 3-5 条)}\n"
)

GREET_PROMPT = (
    "你是求职者本人。基于简历和岗位 JD，写一段在招聘平台（如 BOSS 直聘）发给 HR 的打招呼话术。\n"
    "要求：60-150 字，突出与该岗位最相关的 2-3 个真实亮点，语气真诚自然不浮夸，结尾表达投递意愿。\n"
    "只能基于简历真实内容，禁止编造。直接输出话术正文，不要任何解释。"
)

TAILOR_STREAM_PROMPT = (
    "你是资深简历顾问。基于候选人简历和岗位 JD，生成一份针对该岗位定制的简历。\n"
    "规则：只能改写、重组、强调简历中真实存在的内容，禁止编造经历/技能/数据；"
    "用 STAR 法则改写项目与经历 bullet，把 JD 关键词自然注入真实经历；结构与原简历保持一致。\n"
    "直接输出一份完整、可直接投递的 Markdown 简历，不要 JSON，不要任何多余解释。"
)

TAILOR_PDF_PROMPT = (
    "你是资深简历顾问。基于候选人简历和岗位 JD，生成一份针对该岗位的、能放进一页 A4 的定制简历，输出结构化 JSON。\n"
    "规则：只能改写/重组/强调简历中真实内容，禁止编造；用 STAR 提炼项目要点并自然注入 JD 关键词；"
    "控制篇幅：项目最多 3 个、每个最多 3 条要点、每条不超过 30 字，技能一行。\n"
    "严格只输出 JSON：\n"
    "name: 姓名\n"
    "target: 应聘目标（公司·岗位）\n"
    "contact: 联系方式（电话 · 邮箱 · 城市）\n"
    "education: 教育背景\n"
    "skills: 技能一句话\n"
    "sections: 数组，每项 {title: 板块名, items: [{head: 条目标题(项目名·时间), bullets: [要点]}]}\n"
)


@app.post("/api/resume/tailor")
def tailor_resume(body: dict):
    job, resume = _job_and_resume(body or {})
    d = _llm_json(TAILOR_PROMPT, _mat_user_content(job, resume), kind="定制简历")
    return {
        "resume_md": str(d.get("resume_md", "") or "").strip(),
        "changes": [str(x).strip() for x in (d.get("changes") or []) if str(x).strip()][:6],
    }


@app.post("/api/greeting")
def greeting(body: dict):
    job, resume = _job_and_resume(body or {})
    text = _llm_text(GREET_PROMPT, _mat_user_content(job, resume), kind="打招呼").strip()
    return {"text": text}


@app.get("/api/llm/config")
def llm_config():
    cfg = _llm_cfg()
    try:
        saved = json.loads(_get_setting("saved_models", "[]") or "[]")
    except Exception:
        saved = []
    return {"base": cfg["base"], "key": cfg["key"], "model": cfg["model"], "presets": MODEL_PRESETS, "saved": saved}


@app.post("/api/llm/config")
def llm_set_config(body: dict):
    if "base" in body:
        _set_setting("llm_base", body["base"])
    if "key" in body:
        _set_setting("llm_key", body["key"])
    if "model" in body:
        _set_setting("llm_model", body["model"])
    return {"ok": True}


@app.post("/api/llm/models")
def llm_add_model(body: dict):
    try:
        saved = json.loads(_get_setting("saved_models", "[]") or "[]")
    except Exception:
        saved = []
    saved = [s for s in saved if isinstance(s, dict)]
    item = {
        "name": str(body.get("name") or "模型").strip(),
        "base": str(body.get("base") or "").strip(),
        "key": str(body.get("key") or "").strip(),
        "model": str(body.get("model") or "").strip(),
    }
    saved.append(item)
    _set_setting("saved_models", json.dumps(saved, ensure_ascii=False))
    return {"ok": True, "saved": saved}


@app.delete("/api/llm/models/{idx}")
def llm_del_model(idx: int):
    try:
        saved = json.loads(_get_setting("saved_models", "[]") or "[]")
    except Exception:
        saved = []
    if 0 <= idx < len(saved):
        saved.pop(idx)
    _set_setting("saved_models", json.dumps(saved, ensure_ascii=False))
    return {"ok": True, "saved": saved}


@app.post("/api/llm/test")
def llm_test(body: dict = None):
    body = body or {}
    cfg = _llm_cfg({
        "base": body.get("base") or _llm_cfg()["base"],
        "key": body.get("key") or _llm_cfg()["key"],
        "model": body.get("model") or _llm_cfg()["model"],
    })
    reply = _llm_text("你是一个助手，请回复 OK 两个字。", "hi", override=cfg).strip()
    return {"ok": True, "reply": reply}


def _sse(gen):
    for chunk in gen:
        yield "data: " + json.dumps({"t": chunk}, ensure_ascii=False) + "\n\n"
    yield "data: [DONE]\n\n"


@app.post("/api/tailor/stream")
def tailor_stream(body: dict):
    job, resume = _job_and_resume(body or {})

    def gen():
        try:
            for chunk in _llm_stream(TAILOR_STREAM_PROMPT, _mat_user_content(job, resume), kind="定制简历"):
                yield chunk
        except HTTPException as e:
            yield "\n\n⚠️ " + str(e.detail)

    return StreamingResponse(_sse(gen()), media_type="text/event-stream")


@app.post("/api/greeting/stream")
def greeting_stream(body: dict):
    job, resume = _job_and_resume(body or {})

    def gen():
        try:
            for chunk in _llm_stream(GREET_PROMPT, _mat_user_content(job, resume), kind="打招呼"):
                yield chunk
        except HTTPException as e:
            yield "\n\n⚠️ " + str(e.detail)

    return StreamingResponse(_sse(gen()), media_type="text/event-stream")


@app.post("/api/digest/stream")
def digest_stream(body: dict = None):
    data = _build_digest_data()

    def gen():
        try:
            for chunk in _llm_stream(DIGEST_PROMPT, data, kind="日报"):
                yield chunk
        except HTTPException as e:
            yield "\n\n⚠️ " + str(e.detail)

    return StreamingResponse(_sse(gen()), media_type="text/event-stream")


class ReviewIn(BaseModel):
    company: str = ""
    title: str = ""
    round: str = ""
    date: str = ""
    questions: str = ""
    self_rating: int = 0
    notes: str = ""
    to_jing: bool = True


def _my_interview_folder():
    c = _conn()
    try:
        row = c.execute("SELECT id FROM jing_folders WHERE name='我的面试'").fetchone()
        if row:
            return row["id"]
        with c:
            cur = c.execute("INSERT INTO jing_folders (name,created_at) VALUES (?,?)",
                            ("我的面试", time.strftime("%Y-%m-%d %H:%M")))
        return cur.lastrowid
    finally:
        c.close()


def _mock_folder():
    c = _conn()
    try:
        row = c.execute("SELECT id FROM jing_folders WHERE name='模拟面试'").fetchone()
        if row:
            return row["id"]
        with c:
            cur = c.execute("INSERT INTO jing_folders (name,created_at) VALUES (?,?)",
                            ("模拟面试", time.strftime("%Y-%m-%d %H:%M")))
        return cur.lastrowid
    finally:
        c.close()


def _sync_review_to_jing(rid):
    c = _conn()
    try:
        row = c.execute("SELECT * FROM reviews WHERE id=?", (rid,)).fetchone()
    finally:
        c.close()
    if not row:
        return
    qs = [q.strip() for q in (row["questions"] or "").splitlines() if q.strip()]
    notes = (row["notes"] or "").strip()
    if not qs and not notes:
        return
    content = ("\n".join("Q: " + q for q in qs) + ("\n\n备注：\n" + notes if notes else "")) if qs else notes
    summary = "真实面试复盘 · %s%s" % ((row["round"] or "").strip(), (" " + row["date"]) if (row["date"] or "").strip() else "")
    folder = _my_interview_folder()
    c = _conn()
    try:
        with c:
            if row["jing_id"]:
                c.execute(
                    "UPDATE jings SET company=?, title=?, questions=?, content=?, summary=?, folder_id=? WHERE id=? AND source='review'",
                    (row["company"], row["title"], json.dumps(qs, ensure_ascii=False), content, summary, folder, row["jing_id"]),
                )
                return
            cur = c.execute(
                "INSERT INTO jings (company,title,source_url,content,questions,summary,is_fav,folder_id,source,review_id,created_at) VALUES (?,?,?,?,?,?,?,?,?,?,?)",
                (row["company"], row["title"], "", content, json.dumps(qs, ensure_ascii=False), summary, 0, folder,
                 "review", rid, time.strftime("%Y-%m-%d %H:%M")),
            )
            c.execute("UPDATE reviews SET jing_id=? WHERE id=?", (cur.lastrowid, rid))
    finally:
        c.close()


def _unsync_review_to_jing(rid):
    c = _conn()
    try:
        row = c.execute("SELECT jing_id FROM reviews WHERE id=?", (rid,)).fetchone()
        if row and row["jing_id"]:
            with c:
                c.execute("DELETE FROM jings WHERE id=? AND source='review'", (row["jing_id"],))
                c.execute("UPDATE reviews SET jing_id=0 WHERE id=?", (rid,))
    finally:
        c.close()


REVIEW_PROMPT = (
    "你是资深面试复盘教练。基于候选人记录的面试信息，生成一份复盘报告。\n"
    "内容包含：表现亮点、待改进点、下次面试重点、具体行动清单。\n"
    "用中文，结构化，简洁有力，控制在 400 字内。直接输出报告正文，不要多余解释。"
)


@app.get("/api/reviews")
def list_reviews():
    c = _conn()
    try:
        rows = c.execute("SELECT * FROM reviews ORDER BY id DESC").fetchall()
        return [dict(r) for r in rows]
    finally:
        c.close()


@app.post("/api/reviews")
def add_review(r: ReviewIn):
    c = _conn()
    try:
        with c:
            cur = c.execute(
                "INSERT INTO reviews (company,title,round,date,questions,self_rating,notes,created_at) VALUES (?,?,?,?,?,?,?,?)",
                (r.company, r.title, r.round, r.date, r.questions, r.self_rating, r.notes, time.strftime("%Y-%m-%d %H:%M")),
            )
        rid = cur.lastrowid
    finally:
        c.close()
    if getattr(r, "to_jing", True) and ((r.questions or "").strip() or (r.notes or "").strip()):
        _sync_review_to_jing(rid)
    return {"id": rid}


@app.put("/api/reviews/{rid}")
def update_review(rid: int, patch: dict):
    fields = ("company", "title", "round", "date", "questions", "self_rating", "notes")
    sets, vals = [], []
    for k in fields:
        if k in patch:
            sets.append(k + "=?")
            vals.append(patch[k] if isinstance(patch[k], int) else ("" if patch[k] is None else str(patch[k])))
    c = _conn()
    try:
        cur_row = c.execute("SELECT jing_id FROM reviews WHERE id=?", (rid,)).fetchone()
        if not cur_row:
            raise HTTPException(404, "复盘不存在")
        was_linked = bool(cur_row["jing_id"])
        with c:
            if sets:
                vals.append(rid)
                cur = c.execute("UPDATE reviews SET " + ", ".join(sets) + " WHERE id=?", vals)
                if cur.rowcount == 0:
                    raise HTTPException(404, "复盘不存在")
    finally:
        c.close()
    if "to_jing" in patch:
        if patch["to_jing"]:
            _sync_review_to_jing(rid)
        else:
            _unsync_review_to_jing(rid)
    elif was_linked:
        _sync_review_to_jing(rid)
    return {"ok": True}


@app.delete("/api/reviews/{rid}")
def del_review(rid: int):
    _unsync_review_to_jing(rid)
    c = _conn()
    try:
        with c:
            cur = c.execute("DELETE FROM reviews WHERE id=?", (rid,))
        if cur.rowcount == 0:
            raise HTTPException(404, "复盘不存在")
        return {"ok": True}
    finally:
        c.close()


@app.post("/api/reviews/{rid}/analyze")
def analyze_review(rid: int):
    c = _conn()
    try:
        row = c.execute("SELECT * FROM reviews WHERE id=?", (rid,)).fetchone()
    finally:
        c.close()
    if not row:
        raise HTTPException(404, "复盘不存在")
    if not (row["questions"] or "").strip():
        raise HTTPException(400, "请先填写「被问的问题」再复盘")
    user = ("公司：%s\n岗位：%s\n轮次：%s\n日期：%s\n自评：%s\n被问问题：\n%s\n备注：%s") % (
        row["company"], row["title"], row["round"], row["date"], row["self_rating"] or "未填",
        row["questions"], row["notes"] or "无")
    text = _llm_text(REVIEW_PROMPT, user, kind="复盘").strip()
    c = _conn()
    try:
        with c:
            c.execute("UPDATE reviews SET ai_review=? WHERE id=?", (text, rid))
    finally:
        c.close()
    return {"ai_review": text}


@app.get("/api/cache/stats")
def cache_stats():
    c = _conn()
    try:
        n = c.execute("SELECT COUNT(*) FROM cache").fetchone()[0]
        return {"count": n}
    finally:
        c.close()


@app.delete("/api/cache")
def cache_clear():
    c = _conn()
    try:
        with c:
            c.execute("DELETE FROM cache")
        return {"ok": True}
    finally:
        c.close()


@app.get("/api/usage/summary")
def usage_summary():
    c = _conn()
    try:
        rows = c.execute("SELECT kind, COUNT(*) AS n, SUM(in_chars) AS ic, SUM(out_chars) AS oc, AVG(latency) AS al, SUM(latency) AS tl FROM usage GROUP BY kind").fetchall()
        total = c.execute("SELECT COUNT(*) AS n, SUM(in_chars) AS ic, SUM(out_chars) AS oc, AVG(latency) AS al, SUM(latency) AS tl FROM usage").fetchone()
    finally:
        c.close()
    by_kind = [{"kind": r["kind"], "n": r["n"], "in_chars": r["ic"] or 0, "out_chars": r["oc"] or 0,
                "avg_latency": round(r["al"] or 0, 2), "total_latency": round(r["tl"] or 0, 2)} for r in rows]
    n = total["n"] or 0
    in_chars = total["ic"] or 0
    out_chars = total["oc"] or 0
    in_tokens = int(in_chars / 3)
    out_tokens = int(out_chars / 2)
    cost = round(in_tokens / 1e6 * 0.5 + out_tokens / 1e6 * 2.0, 4)
    return {
        "calls": n, "in_chars": in_chars, "out_chars": out_chars,
        "in_tokens": in_tokens, "out_tokens": out_tokens, "est_cost_yuan": cost,
        "avg_latency": round(total["al"] or 0, 2), "total_latency": round(total["tl"] or 0, 2),
        "by_kind": by_kind,
    }


@app.post("/api/tailor/pdf")
def tailor_pdf(body: dict):
    job, resume = _job_and_resume(body or {})
    d = _llm_json(TAILOR_PDF_PROMPT, _mat_user_content(job, resume), kind="定制简历")
    pdf_bytes = resume_pdf.render_pdf(d)
    fname = (job["company"] or "resume") + "-定制简历.pdf"
    return Response(content=pdf_bytes, media_type="application/pdf",
                    headers={"Content-Disposition": "attachment; filename*=UTF-8''%s" % urllib.parse.quote(fname)})


# ===================== 经历库 & JD 定向改写 =====================

EXP_ITEM_FIELDS = ("type", "company", "title", "role", "time_range", "content")


class ExpItemIn(BaseModel):
    type: str = "实习"
    company: str = ""
    title: str = ""
    role: str = ""
    time_range: str = ""
    content: str = ""


@app.get("/api/exp_items")
def list_exp_items():
    c = _conn()
    try:
        rows = c.execute("SELECT * FROM exp_items ORDER BY id DESC").fetchall()
        return [dict(r) for r in rows]
    finally:
        c.close()


@app.post("/api/exp_items")
def add_exp_item(e: ExpItemIn):
    if e.type not in ("实习", "项目"):
        e.type = "实习"
    if not (e.title or "").strip() and not (e.content or "").strip():
        raise HTTPException(400, "请填写岗位/项目名称或素材内容")
    c = _conn()
    try:
        with c:
            cur = c.execute(
                "INSERT INTO exp_items (type,company,title,role,time_range,content,created_at) VALUES (?,?,?,?,?,?,?)",
                (e.type, e.company, e.title, e.role, e.time_range, e.content, time.strftime("%Y-%m-%d %H:%M")),
            )
        return {"id": cur.lastrowid}
    finally:
        c.close()


@app.put("/api/exp_items/{eid}")
def update_exp_item(eid: int, patch: dict):
    sets, vals = [], []
    for k in EXP_ITEM_FIELDS:
        if k in patch:
            v = patch[k]
            if k == "type" and v not in ("实习", "项目"):
                continue
            sets.append(k + "=?")
            vals.append("" if v is None else str(v))
    if not sets:
        raise HTTPException(400, "无可更新字段")
    vals.append(eid)
    c = _conn()
    try:
        with c:
            cur = c.execute("UPDATE exp_items SET " + ", ".join(sets) + " WHERE id=?", vals)
        if cur.rowcount == 0:
            raise HTTPException(404, "经历不存在")
        return {"ok": True}
    finally:
        c.close()


@app.delete("/api/exp_items/{eid}")
def del_exp_item(eid: int):
    c = _conn()
    try:
        with c:
            cur = c.execute("DELETE FROM exp_items WHERE id=?", (eid,))
        if cur.rowcount == 0:
            raise HTTPException(404, "经历不存在")
        return {"ok": True}
    finally:
        c.close()


def _rule_score(e, req_lines, jd_kws):
    hay = _norm((e.get("content") or "") + (e.get("company") or "") + (e.get("title") or "") + (e.get("role") or ""))
    score = 0
    if req_lines:
        hits = 0
        for line in req_lines:
            if not line:
                continue
            kw = re.findall(r"[a-zA-Z]{2,}|[\u4e00-\u9fa5]{2,}", line)
            if any(k in hay for k in kw):
                hits += 1
        score += min(40, hits * 12)
    if re.search(r"\d+%|\d+万|\d+人|\d+个|\d+次|\d+天", hay):
        score += 15
    m = re.search(r"(20\d{2})", e.get("time_range") or "")
    if m:
        year = int(m.group(1))
        cur = time.localtime().tm_year
        if year >= cur - 1:
            score += 20
        elif year >= cur - 2:
            score += 10
    if jd_kws and any(k in hay for k in jd_kws):
        score += 15
    return min(100, score)


EXP_MATCH_PROMPT = (
    "你是资深简历匹配顾问。下面是一个岗位的 JD 要求，以及若干条候选人的实习/项目经历素材（每条带 ID）。\n"
    "请从中挑选出与该岗位最匹配的 3 条经历，并给出排序与理由。\n"
    "规则：只能从提供的经历里选，禁止编造；优先选择满足 JD 硬性要求、有量化结果、时间更近的经历。\n"
    "严格只输出 JSON：{selected: [{id: 经历ID, reason: 匹配理由(一句话)}]}，最多 3 条。"
)


@app.post("/api/exp/match")
def exp_match(body: dict):
    jid = body.get("job_id")
    if isinstance(jid, str):
        jid = int(jid) if jid.isdigit() else None
    job = _get_job(jid) if jid else None
    if not job:
        raise HTTPException(404, "岗位不存在")
    c = _conn()
    try:
        items = [dict(r) for r in c.execute("SELECT * FROM exp_items ORDER BY id DESC").fetchall()]
    finally:
        c.close()
    if not items:
        return {"candidates": [], "note": "经历库为空，请先在表格页左侧「经历库」添加实习/项目经历"}
    req_lines = [re.sub(r"^[\d\.\-\*\s、]+", "", l).strip() for l in (job["requirements"] or "").splitlines() if l.strip()]
    jd_kws = re.findall(r"[a-zA-Z]{2,}|[\u4e00-\u9fa5]{2,}", (job["title"] or "") + (job["requirements"] or "")[:200])
    jd_kws = [k for k in jd_kws if len(k) >= 2][:20]
    for it in items:
        it["rule_score"] = _rule_score(it, req_lines, jd_kws)
    items.sort(key=lambda x: x["rule_score"], reverse=True)
    top = items[:6]
    selected_ids, reasons = [], {}
    try:
        payload = {
            "job": {"company": job["company"], "title": job["title"], "requirements": job["requirements"]},
            "exps": [{"id": it["id"], "type": it["type"], "company": it["company"], "title": it["title"],
                      "role": it["role"], "time_range": it["time_range"], "content": (it["content"] or "")[:500]}
                     for it in top],
        }
        d = _llm_json(EXP_MATCH_PROMPT, json.dumps(payload, ensure_ascii=False), kind="经历匹配")
        for s in (d.get("selected") or []):
            if isinstance(s, dict) and s.get("id") is not None:
                selected_ids.append(int(s["id"]))
                reasons[int(s["id"])] = str(s.get("reason", "") or "").strip()
    except HTTPException as e:
        log.warning("LLM 精选失败，回退规则打分: %s", e)
        selected_ids = [it["id"] for it in top[:3]]
    selected_ids = selected_ids[:3]
    ordered = [it for it in top if it["id"] in selected_ids] + [it for it in top if it["id"] not in selected_ids]
    ordered = ordered[:6]
    return {
        "candidates": [{
            "id": it["id"], "type": it["type"], "company": it["company"], "title": it["title"],
            "role": it["role"], "time_range": it["time_range"], "content": (it["content"] or "")[:120],
            "rule_score": it["rule_score"], "reason": reasons.get(it["id"], ""),
            "selected": it["id"] in selected_ids,
        } for it in ordered],
        "note": "",
    }


REWRITE_PROMPT = (
    "你是资深简历顾问。基于岗位 JD、用户选定的实习/项目经历素材、以及现用简历（提供姓名/联系方式/教育/技能基线），"
    "生成一份针对该岗位、能放进一页 A4 的定制简历。\n"
    "规则：\n"
    "1. 实习/项目经历只能来自【选定的经历素材】，可润色、重组、强调，但禁止编造数据、技能、奖项\n"
    "2. 用 STAR 提炼要点，自然注入 JD 关键词\n"
    "3. 教育/技能/基本信息沿用现用简历，缺失则省略\n"
    "4. 控制篇幅：实习/项目经历合计最多 3 条，每条最多 3 个要点，每点不超过 30 字；简历完整、可直接投递\n"
    "严格只输出 JSON：\n"
    "name: 姓名\n"
    "target: 应聘目标（公司·岗位）\n"
    "contact: 联系方式（电话 · 邮箱 · 城市）\n"
    "education: 教育背景\n"
    "skills: 技能一句话\n"
    "sections: 数组，每项 {title: 板块名(如 实习经历/项目经历), items: [{head: 条目标题(公司·岗位·时间), bullets: [要点]}]}\n"
    "changes: 数组，针对该岗位做的关键调整 3-5 条"
)


def _clean_resume_json(d):
    if not isinstance(d, dict):
        d = {}
    sections = []
    _BASE_SECTIONS = {"教育背景", "技能", "Education", "Skills"}
    _seen_titles = {}
    for sec in (d.get("sections") or []):
        if not isinstance(sec, dict):
            continue
        title = str(sec.get("title", "") or "").strip()
        if title in _BASE_SECTIONS:
            continue
        items = []
        for it in (sec.get("items") or []):
            if not isinstance(it, dict):
                continue
            bullets = [str(b).strip() for b in (it.get("bullets") or []) if str(b).strip()][:6]
            company = str(it.get("company", "") or "").strip()
            role = str(it.get("role", "") or "").strip()
            tm = str(it.get("time", "") or "").strip()
            head = str(it.get("head", "") or "").strip()
            if not head and (company or role or tm):
                head = " · ".join([x for x in (company, role, tm) if x])
            items.append({"head": head, "company": company, "role": role, "time": tm, "bullets": bullets})
        if items:
            if title in _seen_titles:
                # 同板块重复出现：保留条目多的那个
                idx = _seen_titles[title]
                if len(items) > len(sections[idx]["items"]):
                    sections[idx] = {"title": title, "items": items}
            else:
                _seen_titles[title] = len(sections)
                sections.append({"title": title, "items": items})
    return {
        "name": str(d.get("name", "") or "").strip(),
        "target": str(d.get("target", "") or "").strip(),
        "contact": str(d.get("contact", "") or "").strip(),
        "education": str(d.get("education", "") or "").strip(),
        "skills": str(d.get("skills", "") or "").strip(),
        "sections": sections[:12],
        "changes": [str(x).strip() for x in (d.get("changes") or []) if str(x).strip()][:6],
    }


def _md_from_resume(d, lang="zh"):
    edu_t = "Education" if lang == "en" else "教育背景"
    skill_t = "Skills" if lang == "en" else "技能"
    lines = []
    name = str(d.get("name", "") or "").strip()
    target = str(d.get("target", "") or "").strip()
    contact = str(d.get("contact", "") or "").strip()
    if name:
        lines.append("# " + name)
    if target:
        lines.append(target)
    if contact:
        lines.append(contact)
    if name or target or contact:
        lines.append("")
    if d.get("education"):
        lines += ["## " + edu_t, str(d["education"]), ""]
    if d.get("skills"):
        lines += ["## " + skill_t, str(d["skills"]), ""]
    for sec in d.get("sections") or []:
        title = str(sec.get("title", "") or "经历").strip()
        items = sec.get("items") or []
        if not items:
            continue
        lines.append("## " + title)
        for it in items:
            head = str(it.get("head", "") or "").strip()
            if head:
                lines.append("### " + head)
            for b in it.get("bullets") or []:
                b = str(b).strip()
                if b:
                    lines.append("- " + b)
        lines.append("")
    return "\n".join(lines).strip() + "\n"


def _resume_from_md(md):
    """把应用生成的 Markdown 简历解析回结构化 JSON（供用户编辑后重新生成 PDF/翻译）。"""
    if not (md or "").strip():
        return None
    res = {"name": "", "target": "", "contact": "", "education": "", "skills": "", "sections": [], "changes": []}
    section = None
    item = None
    pending = None
    for raw in (md or "").splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        if line.startswith("# "):
            res["name"] = line[2:].strip()
        elif line.startswith("## "):
            t = line[3:].strip()
            if t in ("教育背景", "Education"):
                pending = "education"; section = None; item = None
            elif t in ("技能", "Skills"):
                pending = "skills"; section = None; item = None
            else:
                pending = None
                section = {"title": t, "items": []}
                res["sections"].append(section)
                item = None
        elif line.startswith("### "):
            head = line[4:].strip()
            if section is None:
                section = {"title": "经历", "items": []}
                res["sections"].append(section)
            item = {"head": head, "bullets": [], "company": "", "role": "", "time": ""}
            parts = [p.strip() for p in head.split(" · ")]
            if len(parts) >= 3:
                item["company"], item["role"], item["time"] = parts[0], parts[1], " · ".join(parts[2:])
            elif len(parts) == 2:
                item["company"], item["role"], item["time"] = parts[0], parts[1], ""
            elif parts:
                item["company"] = parts[0]
            section["items"].append(item)
        elif line.startswith("- "):
            b = line[2:].strip()
            if item is not None:
                item["bullets"].append(b)
            elif section is not None:
                if not section["items"]:
                    section["items"].append({"head": "", "bullets": [], "company": "", "role": "", "time": ""})
                    item = section["items"][-1]
                item["bullets"].append(b)
        else:
            txt = line.strip()
            if pending == "education":
                res["education"] = (res["education"] + (" " if res["education"] else "") + txt).strip()
            elif pending == "skills":
                res["skills"] = (res["skills"] + (" " if res["skills"] else "") + txt).strip()
            elif section is None and pending is None and res["name"]:
                if not res["target"]:
                    res["target"] = txt
                elif not res["contact"]:
                    res["contact"] = txt
    return res


@app.post("/api/resume/rewrite")
def resume_rewrite(body: dict):
    job, resume = _job_and_resume(body or {})
    exp_ids = body.get("exp_ids") or []
    try:
        exp_ids = [int(x) for x in exp_ids]
    except Exception:
        exp_ids = []
    c = _conn()
    try:
        if exp_ids:
            ph = ",".join("?" * len(exp_ids))
            rows = c.execute("SELECT * FROM exp_items WHERE id IN (%s) ORDER BY id DESC" % ph, exp_ids).fetchall()
        else:
            rows = []
    finally:
        c.close()
    if not rows:
        raise HTTPException(400, "请先在左侧「经历库」勾选至少 1 条实习/项目经历")
    exps_blk = []
    for r in rows:
        exps_blk.append("[%s] %s · %s · %s (%s)\n素材：%s" % (
            r["type"], r["company"] or "?", r["title"] or "?", r["role"] or "?", r["time_range"] or "?",
            r["content"] or "（无素材）"))
    user = ("【岗位信息】\n公司：%s\n岗位：%s\nBase：%s\n要求：\n%s\n\n【JD 原文】\n%s\n\n"
            "【现用简历（提供姓名/联系方式/教育/技能基线）】\n%s\n\n"
            "【用户选定的实习/项目经历素材（经历只能从中取材）】\n%s") % (
        job["company"], job["title"], job["base"], job["requirements"] or "（无）",
        (job["jd_text"] or "（无）")[:3000], (resume["content"] or "")[:4000],
        "\n\n".join(exps_blk),
    )
    d = _llm_json(REWRITE_PROMPT, user, kind="定制简历")
    d = _clean_resume_json(d)
    return {"resume": d, "resume_md": _md_from_resume(d)}


# ===================== 评估 + 改写（分组 prompt） =====================

EVAL_PROMPT = (
    "你是校招简历智能评估引擎。基于岗位JD、候选人简历和用户补充的真实经历，输出结构化评估报告，并推荐最适合改写的经历。严格遵守规则，违反约束的输出视为无效。\n"
    "输入：\n"
    "<jd>岗位JD（公司/岗位/要求/JD原文）</jd>\n"
    "<resume>候选人简历全文</resume>\n"
    "<extra_experiences>用户补充经历（可能为空）</extra_experiences>\n"
    "<target_format>一页、单栏、标准字体、无表格图片、每条经历2-4个bullet、ATS可解析</target_format>\n"
    "严格按顺序执行：\n"
    "1) JD解析：提取 job_title；hard_requirements(education/major/graduation_year/certificates/other_mandatory)；"
    "must_have_skills[{skill,category,evidence_hint}]；nice_to_have_skills；experience_requirements[{type,level,content}]；achievement_keywords；responsibility_keywords。\n"
    "2) 简历解析：合并简历与补充经历，提取 basic_info(education_level/school/major/graduation_year/gpa) 与 experiences 列表"
    "（每条含 id/type/company_or_project/duration/role/responsibilities/tech_stack/achievements/source）。\n"
    "3) 否决判定：仅对 JD 中**明确强制**的条件做否决（如学历、专业、毕业年份、明确要求且必须提供的证书、其他\"必须\"条件）。"
    "注意：① 语言类要求做**等价换算**——雅思 6.5 及以上 ≥ CET-6、雅思 6.0 ≥ CET-4，托福/专四专八/六级等任意更高或等价证明均视为满足，"
    "不得因简历未写\"CET\"字样而否决英语要求；"
    "② 简历未提及某项条件时标记为\"信息不足\"并视为通过（passed=true），不触发否决；"
    "③ 技能类、经验类、偏好类要求一律不作为否决项。"
    "输出 veto_triggered 与 veto_details[{item,jd_requirement,candidate_status,passed}]；"
    "确认触发时 final_score=0 并跳过第4步评分，但仍做第5步经历推荐。\n"
    "4) 分层评分：总分=核心匹配度×70%+潜力加分×20%+格式可读性×10%。"
    "核心匹配度：技能匹配(55%)双通道——精确关键词命中率 + 语义等价匹配(计0.7，需输出简历原文证据)；"
    "经历适配(45%)=经历层级(40%；大厂核心100/大厂边缘或中厂核心75/中厂普通或小厂核心50/无0) + 经历内容匹配(60%；对每条experience_requirement找最相似经历，输出similarity与resume_quote)。"
    "潜力加分：教育背景质量、成就量化密度、差异化亮点 三均值。"
    "格式可读性：单栏/标准字体/无表格图片/信息层级清晰/ATS可解析，每项问题扣20分，最低0分。"
    "修正机制：技能匹配<40→min(raw,50)；经历内容匹配=0→min(raw,45)；经历层级=0且内容<30→min(raw,40)。"
    "映射：75-100推荐 / 55-74待定 / 0-54不推荐。评分必须输出证据链与弱点。\n"
    "5) 经历推荐：从 experiences 中推荐与该JD最相关的 2-3 条（recommended_exp_ids），优先：must_have_skills重合高 > 有量化成果 > 平台层级高。\n"
    "只输出合法JSON（无任何其他文字，用中文）：\n"
    "{\"basic_info\":{\"name\":\"\",\"target\":\"\",\"contact\":\"\",\"education\":\"\",\"skills\":\"\"},\n"
    " \"evaluation\":{\"veto_triggered\":false,\"veto_details\":[],\"final_score\":0,\"recommendation\":\"推荐\",\n"
    "   \"dimension_scores\":{\"skill_match\":{\"score\":0,\"exact_hits\":[],\"semantic_hits\":[],\"misses\":[]},\n"
    "     \"experience_match\":{\"score\":0,\"level_score\":0,\"content_score\":0,\"evidence\":[]},\n"
    "     \"potential_score\":0,\"format_score\":0},\n"
    "   \"evidence_chain\":[],\"weaknesses\":[],\"suggestions\":[]},\n"
    " \"recommended_exp_ids\":[]}\n"
    "约束：事实优先禁止编造；评分必须有证据；同一问题不重复扣分；信息不足输出\"信息不足，建议用户补充\"。"
)


def _eval_cfg():
    base = _get_setting("eval_base", "").strip() or _llm_cfg()["base"]
    key = _get_setting("eval_key", "").strip() or _llm_cfg()["key"]
    model = _get_setting("eval_model", "").strip()
    if not model and "opencode.ai" in base:
        model = "deepseek-v4-flash"
    return {"base": base.rstrip("/"), "key": key, "model": model}


@app.get("/api/eval/config")
def eval_config():
    return {
        "model": _get_setting("eval_model", ""),
        "base": _get_setting("eval_base", ""),
        "key": _get_setting("eval_key", ""),
        "effective": _eval_cfg(),
    }


@app.post("/api/eval/config")
def eval_set_config(body: dict):
    for k, v in body.items():
        if k in ("eval_model", "eval_base", "eval_key"):
            _set_setting(k, str(v or "").strip())
    return {"ok": True}


def _evaluate_user_content(job, resume, exps):
    exp_blk = []
    for i, e in enumerate(exps, 1):
        exp_blk.append("[%d] %s | %s | %s | %s | %s\n   描述：%s" % (
            i, e["type"], e["company"] or "?", e["title"] or "?", e["role"] or "?", e["time_range"] or "?", e["content"] or "（无）"))
    jd_blk = "公司：%s\n岗位：%s\nBase：%s\n要求：\n%s\nJD原文：\n%s" % (
        job["company"], job["title"], job["base"], job["requirements"] or "（无）", (job["jd_text"] or "")[:3000])
    return ("<jd>\n%s\n</jd>\n\n<resume>\n%s\n</resume>\n\n<extra_experiences>\n%s\n</extra_experiences>\n\n"
            "<target_format>一页、单栏、标准字体、无表格图片、每条经历2-4个bullet、ATS可解析</target_format>") % (
        jd_blk, (resume["content"] or "")[:6000], "\n".join(exp_blk) if exp_blk else "（无）")


def _run_evaluation(job, resume):
    """统一的「JD 评估」引擎：解析 + 否决 + 分层评分 + 经历推荐。供评估接口与一键匹配共用，保证分数一致。"""
    c = _conn()
    try:
        exps = [dict(r) for r in c.execute("SELECT * FROM exp_items ORDER BY id DESC").fetchall()]
    finally:
        c.close()
    d = _llm_json(EVAL_PROMPT, _evaluate_user_content(job, resume, exps), override=_eval_cfg(), kind="评估")
    evaluation = d.get("evaluation") or {}
    if not isinstance(evaluation, dict):
        evaluation = {}
    basic = d.get("basic_info") or {}
    if not isinstance(basic, dict):
        basic = {}
    rec_ids = set()
    for x in (d.get("recommended_exp_ids") or []):
        try:
            rec_ids.add(int(x))
        except Exception:
            pass
    req_lines = [re.sub(r"^[\d\.\-\*\s、]+", "", l).strip() for l in (job["requirements"] or "").splitlines() if l.strip()]
    jd_kws = re.findall(r"[a-zA-Z]{2,}|[\u4e00-\u9fa5]{2,}", (job["title"] or "") + (job["requirements"] or "")[:200])
    jd_kws = [k for k in jd_kws if len(k) >= 2][:20]
    if not rec_ids and exps:
        ranked = sorted(exps, key=lambda e: _rule_score(e, req_lines, jd_kws), reverse=True)
        rec_ids = {e["id"] for e in ranked[:3]}
    candidates = []
    for e in exps:
        candidates.append({
            "id": e["id"], "type": e["type"], "company": e["company"], "title": e["title"],
            "role": e["role"], "time_range": e["time_range"], "content": (e["content"] or "")[:120],
            "rule_score": _rule_score(e, req_lines, jd_kws), "reason": "",
            "selected": e["id"] in rec_ids,
        })
    candidates.sort(key=lambda x: (x["selected"], x["rule_score"]), reverse=True)
    return {
        "basic_info": {"name": str(basic.get("name", "") or "").strip(), "target": str(basic.get("target", "") or "").strip(),
                       "contact": str(basic.get("contact", "") or "").strip(), "education": str(basic.get("education", "") or "").strip(),
                       "skills": str(basic.get("skills", "") or "").strip()},
        "evaluation": evaluation,
        "candidates": candidates[:10],
        "note": "" if exps else "经历库为空，评估仅基于现有简历；可在表格页左侧「经历库」补充经历后重新评估",
    }


@app.post("/api/resume/evaluate")
def resume_evaluate(body: dict):
    job, resume = _job_and_resume(body or {})
    return _run_evaluation(job, resume)


POSITION_DIMS = [
    {
        "name": "产品/AI产品",
        "keys": ["产品经理", "AI产品", "产品运营", "产品", "需求"],
        "guide": "突出产品方法论、需求分析与PRD、原型设计、用户研究、数据驱动决策、跨团队协作；有AI/大模型相关经历重点呈现",
    },
    {
        "name": "数据/分析",
        "keys": ["数据产品", "数据分析", "数据运营", "商业分析", "BI", "数据开发"],
        "guide": "突出数据分析、SQL/Python取数与建模、指标体系搭建、可视化看板、A/B测试、数据驱动决策",
    },
    {
        "name": "运营",
        "keys": ["运营", "用户运营", "内容运营", "增长", "活动策划", "私域"],
        "guide": "突出增长方法论、活动策划与执行、用户洞察、数据分析、内容/渠道运营、转化与留存指标",
    },
    {
        "name": "市场/商务/GTM",
        "keys": ["市场", "商务", "GTM", "BD", "渠道", "销售", "推广"],
        "guide": "突出市场洞察、渠道拓展、商务沟通、方案撰写、数据归因、客户/商户需求理解",
    },
    {
        "name": "财务/审计/咨询",
        "keys": ["财务", "审计", "咨询", "分析师", "会计", "投资", "证券"],
        "guide": "突出财务/业务分析、建模估值、尽调、报告撰写、跨部门协调、专业工具使用",
    },
    {
        "name": "技术开发",
        "keys": ["开发", "工程师", "后端", "前端", "算法", "测试", "全栈", "技术"],
        "guide": "突出编程语言、框架、数据库、系统设计、工程实践、性能优化、代码质量",
    },
]


def _position_guide(title="", requirements=""):
    text = (title or "") + " " + (requirements or "")[:200]
    for d in POSITION_DIMS:
        if any(k in text for k in d["keys"]):
            return d["guide"]
    return ""


REWRITE2_PROMPT = (
    "你是校招简历改写引擎。你的任务是基于岗位JD、评估结果、用户选定的经历素材，对候选人原始简历做「丰富 + 润色」，"
    "保持原有格式与完整性，输出一页 A4 满满当当的定制简历。严格遵守规则，违反约束的输出视为无效。\n"
    "规则：\n"
    "1. 保留原始简历的完整板块结构与顺序（如 教育背景/实习经历/项目经历/科研经历/竞赛经历/技能/获奖 等），"
    "禁止合并、删除、篡改板块；只允许在合适位置补充板块（如新增的获奖/技能）\n"
    "2. 内容必须覆盖简历中全部真实经历与信息，禁止编造技术/工具/职责/数字/成果；禁止把\"参与\"改写为\"主导\"、\"协助\"改写为\"负责\"\n"
    "3. 每条要点格式固定为「动作小标题：背景说明，用什么工具做了什么，达成了什么成果」。"
    "小标题用 2-6 字的关键动作词（如：看板搭建、A/B测试、需求调研、模型调优）；"
    "冒号后先写背景/职责，再写所用工具或方法，最后写量化成果（数字必须来自原文）；每段经历 3-5 条要点，自然嵌入JD关键词（不重复堆砌）\n"
    "4. 依据评估结果调整优先级：与JD最匹配的经历可适当前置或略多写，但不得因此删减其他真实经历\n"
    "5. 教育/技能板块优先沿用简历原文的完整内容，可润色，不要精简丢失\n"
    "6. 一页 A4 为目标：通过简洁有力的 bullet 语言把内容写满一页，而不是删减内容\n"
    "7. 无表格、无图片、ATS 可解析\n"
    "8. 若评估 veto_triggered=true：仅做语言润色，不调整经历事实\n"
    "9. 若输入含 <position_guide>，按该岗位类型侧重的核心能力维度调整内容侧重（通过措辞与要点排序体现，不删减其他真实经历）\n"
    "严格只输出 JSON：\n"
    "name: 姓名\n"
    "target: 应聘目标（公司·岗位）\n"
    "contact: 联系方式\n"
    "education: 教育背景（完整）\n"
    "skills: 技能（完整）\n"
    "sections: 数组，每项 {title: 板块名(中文), items: [{company: 公司/机构名, role: 岗位/角色, time: 时间, bullets: [要点]}]}，板块顺序与原始简历一致，板块数不限\n"
    "changes: 数组，本次改写的关键调整 3-5 条\n"
    "只输出 JSON，不要任何其他文字。"
)


@app.post("/api/resume/rewrite_v2")
def resume_rewrite_v2(body: dict):
    job, resume = _job_and_resume(body or {})
    exp_ids = body.get("exp_ids") or []
    try:
        exp_ids = [int(x) for x in exp_ids]
    except Exception:
        exp_ids = []
    c = _conn()
    try:
        if exp_ids:
            ph = ",".join("?" * len(exp_ids))
            rows = c.execute("SELECT * FROM exp_items WHERE id IN (%s) ORDER BY id DESC" % ph, exp_ids).fetchall()
        else:
            rows = []
    finally:
        c.close()
    exp_blk = []
    for r in rows:
        exp_blk.append("[%s] %s · %s · %s (%s)\n素材：%s" % (
            r["type"], r["company"] or "?", r["title"] or "?", r["role"] or "?", r["time_range"] or "?", r["content"] or "（无素材）"))
    evaluation = body.get("evaluation") or {}
    basic = body.get("basic_info") or {}
    if not isinstance(evaluation, dict):
        evaluation = {}
    if not isinstance(basic, dict):
        basic = {}
    user = ("<jd>\n公司：%s\n岗位：%s\nBase：%s\n要求：\n%s\nJD原文：\n%s\n</jd>\n\n"
            "<evaluation>\n%s\n</evaluation>\n\n"
            "<basic_info>\n%s\n</basic_info>\n\n"
            "<selected_experiences>\n%s\n</selected_experiences>\n\n"
            "<resume_original>\n%s\n</resume_original>") % (
        job["company"], job["title"], job["base"], job["requirements"] or "（无）", (job["jd_text"] or "")[:3000],
        json.dumps(evaluation, ensure_ascii=False)[:5000],
        json.dumps(basic, ensure_ascii=False),
        "\n".join(exp_blk) if exp_blk else "（无选定经历，使用 resume_original 中的经历，同样禁止编造）",
        (resume["content"] or "")[:4000],
    )
    guide = _position_guide(job["title"], job["requirements"])
    if guide:
        user += "\n\n<position_guide>%s</position_guide>" % guide
    d = _llm_json(REWRITE2_PROMPT, user, override=_eval_cfg(), kind="定制简历")
    d = _clean_resume_json(d)
    return {"resume": d, "resume_md": _md_from_resume(d)}


@app.post("/api/resume/md_to_structure")
def resume_md_to_structure(body: dict):
    content = (body.get("content") or "").strip()
    if len(content) < 20:
        raise HTTPException(400, "内容太短")
    d = _resume_from_md(content)
    if not d or not d.get("sections"):
        raise HTTPException(400, "无法解析：请保持 markdown 结构（## 板块 / ### 条目 / - 要点）")
    d = _clean_resume_json(d)
    return {"resume": d, "resume_md": _md_from_resume(d)}


POLISH_PROMPT = (
    "你是资深中文简历润色专家。对这份结构化简历做语言润色，让表达更专业、圆润、有力，提升中文库表达质量。\n"
    "规则：\n"
    "1. 保持 JSON 结构完全一致（字段名/板块/条目/时间/顺序不变），只润色文字\n"
    "2. 保留全部事实、公司名、岗位、时间与数字，禁止编造/夸大/增删经历\n"
    "3. 每条要点保持「动作小标题：背景说明，用什么工具做了什么，达成什么成果」格式\n"
    "4. 用强动词（搭建/优化/推动/主导/落地）、成果导向、STAR 思路；减少空洞形容词（如\"很好\"\"负责相关\"）；表达凝练圆润、有力度\n"
    "5. 用中文\n"
    "只输出润色后的 JSON，不要任何其他文字。"
)


@app.post("/api/resume/polish")
def resume_polish(body: dict):
    resume = body.get("resume") or {}
    if not isinstance(resume, dict) or not resume.get("sections"):
        raise HTTPException(400, "缺少简历数据，请先生成改写简历")
    prompt = POLISH_PROMPT
    guide = _position_guide(str(body.get("job_title") or ""))
    if guide:
        prompt += "\n岗位类型侧重：" + guide
    d = _llm_json(prompt, json.dumps(resume, ensure_ascii=False), override=_eval_cfg(), kind="润色")
    d = _clean_resume_json(d)
    return {"resume": d, "resume_md": _md_from_resume(d)}


TRANSLATE_PROMPT = (
    "你是专业英文简历翻译专家。把下面这份结构化中文简历翻译成地道、正式的商务英语简历。\n"
    "规则：\n"
    "1. 保持 JSON 结构完全一致（字段名不变），只翻译字段的值\n"
    "2. 商务英语，术语专业（实习经历=Internship Experience，项目经历=Project Experience，教育背景=Education，技能=Skills）\n"
    "3. 人名用拼音；公司/机构名优先用官方英文名（如 字节跳动=ByteDance），否则用拼音\n"
    "4. 时间格式不变；每条要点保持简洁（不超过 15 个英文单词）\n"
    "5. 禁止编造任何中文原文中没有的内容\n"
    "严格只输出翻译后的 JSON。"
)


def _translate_deepl_texts(texts, target="EN-US"):
    key = _get_setting("deepl_key", "").strip()
    if not key:
        raise RuntimeError("未配置 DeepL API Key")
    host = "https://api-free.deepl.com" if key.endswith(":fx") else "https://api.deepl.com"
    out = []
    for i in range(0, len(texts), 40):
        chunk = texts[i:i + 40]
        r = requests.post(host + "/v2/translate",
                          json={"text": chunk, "target_lang": target},
                          headers={"Authorization": "DeepL-Auth-Key " + key},
                          timeout=60)
        d = r.json()
        if "translations" not in d:
            raise RuntimeError("DeepL 调用失败: %s" % str(d)[:200])
        out += [t["text"] for t in d["translations"]]
    return out


def _translate_deepl_resume(resume):
    import copy
    en = copy.deepcopy(resume)
    texts = []

    def collect(o):
        if isinstance(o, dict):
            for v in o.values():
                collect(v)
        elif isinstance(o, list):
            for it in o:
                collect(it)
        else:
            texts.append(str(o))

    collect(resume)
    if not texts:
        return en
    translated = _translate_deepl_texts(texts)
    idx = [0]

    def fill(o):
        if isinstance(o, dict):
            for k, v in o.items():
                if isinstance(v, str):
                    o[k] = translated[idx[0]]
                    idx[0] += 1
                else:
                    fill(v)
        elif isinstance(o, list):
            for it in o:
                fill(it)

    fill(en)
    return en


@app.post("/api/resume/translate")
def resume_translate(body: dict):
    resume = body.get("resume") or {}
    if not isinstance(resume, dict) or not resume.get("sections"):
        raise HTTPException(400, "缺少简历数据，请先生成定制简历")
    provider = _get_setting("translate_provider", "llm")
    try:
        if provider == "deepl":
            en = _translate_deepl_resume(resume)
        else:
            en = _llm_json(TRANSLATE_PROMPT, json.dumps(resume, ensure_ascii=False), override=_eval_cfg(), kind="翻译")
    except Exception as e:
        raise HTTPException(502, "翻译失败：%s" % e)
    en = _clean_resume_json(en)
    return {"resume_en": en, "resume_en_md": _md_from_resume(en, lang="en")}


TRANSLATE_TEXT_PROMPT = (
    "你是资深英文简历翻译专家。把下面这份中文简历翻译成地道、正式的商务英语简历，并结构化输出。\n"
    "规则：\n"
    "1. 内容只能来自原文，禁止编造经历/技能/数据/奖项\n"
    "2. 保留原文的板块与顺序（教育、实习经历、项目经历、技能、获奖、证书等），信息完整不遗漏\n"
    "3. 商务英语，术语专业；人名用拼音；公司/机构名优先用官方英文名（如 字节跳动=ByteDance），否则用拼音\n"
    "4. 每个经历/项目的要点 3-5 条，每条不超过 15 个英文单词，简洁有力，突出量化结果\n"
    "5. 原文若为 Markdown，忽略格式符号，只保留内容\n"
    "严格只输出 JSON：\n"
    "name: 姓名\n"
    "target: 求职意向/应聘目标（没有则空字符串）\n"
    "contact: 联系方式（电话 · 邮箱 · 城市）\n"
    "education: 教育背景\n"
    "skills: 技能一句话\n"
    "sections: 数组，每项 {title: 板块名(英文), items: [{company: 公司/机构名(英文), role: 岗位/角色(英文), time: 时间, bullets: [要点]}]}\n"
    "若某板块（如获奖/证书）不适合放 education/skills，请放入 sections 的相应板块。"
)


@app.post("/api/resume/translate_text")
def resume_translate_text(body: dict):
    content = (body.get("content") or "").strip()
    if len(content) < 20:
        raise HTTPException(400, "简历内容太短")
    try:
        en = _llm_json(TRANSLATE_TEXT_PROMPT, content[:8000], override=_eval_cfg(), kind="翻译")
    except Exception as e:
        raise HTTPException(502, "翻译失败：%s" % e)
    en = _clean_resume_json(en)
    return {"resume_en": en, "resume_en_md": _md_from_resume(en, lang="en")}


@app.post("/api/resume/pdf")
def resume_pdf_endpoint(body: dict):
    resume = body.get("resume") or {}
    lang = str(body.get("lang") or "zh").strip()
    if lang not in ("zh", "en"):
        lang = "zh"
    if not isinstance(resume, dict) or not resume.get("sections"):
        raise HTTPException(400, "缺少简历数据")
    photo = str(PHOTO_PATH) if PHOTO_PATH.exists() else None
    pdf_bytes = resume_pdf.render_pdf(resume, lang=lang, photo=photo)
    fname = "定制简历-" + ("EN" if lang == "en" else "中文") + ".pdf"
    return Response(content=pdf_bytes, media_type="application/pdf",
                    headers={"Content-Disposition": "attachment; filename*=UTF-8''%s" % urllib.parse.quote(fname)})


@app.get("/api/translate/config")
def translate_config():
    return {"provider": _get_setting("translate_provider", "llm"), "deepl_key": _get_setting("deepl_key", "")}


@app.post("/api/translate/config")
def translate_set_config(body: dict):
    if "provider" in body:
        v = str(body.get("provider") or "llm").strip()
        if v not in ("llm", "deepl"):
            raise HTTPException(400, "非法翻译方式")
        _set_setting("translate_provider", v)
    if "deepl_key" in body:
        _set_setting("deepl_key", str(body.get("deepl_key") or "").strip())
    return {"ok": True}


@app.get("/api/stages")
def list_stages():
    return {"stages": _all_stages()}


@app.post("/api/stages")
def add_stage(body: dict):
    name = str(body.get("name") or "").strip()
    if not name:
        raise HTTPException(400, "名称不能为空")
    if name in _all_stages():
        raise HTTPException(400, "该阶段已存在")
    try:
        custom = json.loads(_get_setting("custom_stages", "[]") or "[]")
        custom = [str(x) for x in custom if str(x).strip()]
    except Exception:
        custom = []
    custom.append(name)
    _set_setting("custom_stages", json.dumps(custom, ensure_ascii=False))
    return {"ok": True, "stages": _all_stages()}


@app.delete("/api/stages/{name}")
def del_stage(name: str):
    try:
        custom = json.loads(_get_setting("custom_stages", "[]") or "[]")
    except Exception:
        custom = []
    custom = [x for x in custom if x != name]
    _set_setting("custom_stages", json.dumps(custom, ensure_ascii=False))
    return {"ok": True, "stages": _all_stages()}


JING_PARSE_PROMPT = (
    "你是面经整理助手。从候选人提供的面经原文中提取信息，输出结构化 JSON：\n"
    "company: 公司名称\n"
    "title: 岗位\n"
    "questions: 数组，逐条列出面试中被问到的真实问题（保留原文，最多 15 条）\n"
    "summary: 2-3 句要点总结（难度、流程、注意事项）\n"
    "原文中不存在的信息不要编造，直接省略对应字段。"
)


@app.post("/api/jing/parse")
def jing_parse(body: dict):
    text = (body.get("text") or "").strip()
    if len(text) < 20:
        raise HTTPException(400, "面经内容太短")
    d = _llm_json(JING_PARSE_PROMPT, text[:6000], kind="面经")
    return {
        "company": str(d.get("company", "") or "").strip(),
        "title": str(d.get("title", "") or "").strip(),
        "questions": [str(q).strip() for q in (d.get("questions") or []) if str(q).strip()][:15],
        "summary": str(d.get("summary", "") or "").strip(),
    }


@app.post("/api/jing/add")
def jing_add(body: dict):
    company = str(body.get("company") or "").strip()
    if not company:
        raise HTTPException(400, "请填写公司名称")
    title = str(body.get("title") or "").strip()
    source_url = str(body.get("source_url") or "").strip()
    content = str(body.get("content") or "").strip()
    questions = [str(q) for q in (body.get("questions") or []) if str(q).strip()]
    summary = str(body.get("summary") or "").strip()
    folder_id = int(body.get("folder_id") or 0)
    if not content and not questions:
        raise HTTPException(400, "内容不能为空")
    if not questions and content:
        try:
            d = _llm_json(JING_PARSE_PROMPT, content[:6000], kind="面经")
            questions = [str(q).strip() for q in (d.get("questions") or []) if str(q).strip()][:15]
            if not summary:
                summary = str(d.get("summary", "") or "").strip()
        except Exception:
            pass
    c = _conn()
    try:
        with c:
            cur = c.execute(
                "INSERT INTO jings (company,title,source_url,content,questions,summary,is_fav,folder_id,created_at) VALUES (?,?,?,?,?,?,?,?,?)",
                (company, title, source_url, content, json.dumps(questions, ensure_ascii=False), summary, 0, folder_id, time.strftime("%Y-%m-%d %H:%M")),
            )
        return {"id": cur.lastrowid}
    finally:
        c.close()


@app.get("/api/jings")
def list_jings(company: str = "", folder: int = 0, q: str = ""):
    sql = "SELECT * FROM jings WHERE 1=1"
    args = []
    if company:
        sql += " AND company=?"
        args.append(company)
    if folder:
        sql += " AND folder_id=?"
        args.append(folder)
    if q:
        sql += " AND (company LIKE ? OR title LIKE ? OR content LIKE ? OR questions LIKE ?)"
        like = "%" + q + "%"
        args += [like, like, like, like]
    sql += " ORDER BY is_fav DESC, id DESC"
    c = _conn()
    try:
        rows = c.execute(sql, args).fetchall()
        out = []
        for r in rows:
            d = dict(r)
            try:
                d["questions"] = json.loads(r["questions"] or "[]")
            except Exception:
                d["questions"] = []
            out.append(d)
        return out
    finally:
        c.close()


@app.put("/api/jings/{jid}")
def update_jing(jid: int, patch: dict):
    sets, vals = [], []
    for k in ("is_fav", "title", "summary", "folder_id"):
        if k in patch:
            v = patch[k]
            if k == "is_fav":
                v = 1 if v else 0
            elif k == "folder_id":
                v = int(v or 0)
            else:
                v = "" if v is None else str(v)
            sets.append(k + "=?")
            vals.append(v)
    if not sets:
        raise HTTPException(400, "无可更新字段")
    vals.append(jid)
    c = _conn()
    try:
        with c:
            cur = c.execute("UPDATE jings SET " + ", ".join(sets) + " WHERE id=?", vals)
        if cur.rowcount == 0:
            raise HTTPException(404, "面经不存在")
        return {"ok": True}
    finally:
        c.close()


@app.delete("/api/jings/{jid}")
def del_jing(jid: int):
    c = _conn()
    try:
        with c:
            cur = c.execute("DELETE FROM jings WHERE id=?", (jid,))
        if cur.rowcount == 0:
            raise HTTPException(404, "面经不存在")
        return {"ok": True}
    finally:
        c.close()


def _open_url(url, headers, proxy=None, timeout=30):
    import urllib.request

    if proxy:
        opener = urllib.request.build_opener(urllib.request.ProxyHandler({"http": proxy, "https": proxy}))
    else:
        opener = urllib.request.build_opener()
    req = urllib.request.Request(url, headers=headers)
    with opener.open(req, timeout=timeout) as resp:
        return resp.read().decode("utf-8", errors="ignore")


def _fetch_direct(url, headers, proxy=None):
    from html import unescape

    fetch_url = url
    if "reddit.com" in url and not url.endswith(".json"):
        fetch_url = url.rstrip("/") + ".json"
    raw = _open_url(fetch_url, headers, proxy, timeout=8)
    if fetch_url.endswith(".json") or raw.lstrip().startswith("["):
        data = json.loads(raw)
        try:
            post = data[0]["data"]["children"][0]["data"]
            return (post.get("selftext") or "") + "\n" + (post.get("title") or "")
        except Exception:
            return raw
    body = re.search(r"<body.*?</body>", raw, re.S)
    if body:
        t = re.sub(r"<script.*?</script>|<style.*?</style>", "", body.group(0), flags=re.S)
        t = re.sub(r"<[^>]+>", "\n", t)
        t = unescape(t)
        return "\n".join(l.strip() for l in t.splitlines() if l.strip())
    return raw[:8000]


def _fetch_jina(url, proxy=None):
    return _open_url("https://r.jina.ai/" + url, {"User-Agent": "Mozilla/5.0"}, proxy, timeout=15)


# ===================== 小红书抓取 & 视觉转写 =====================

def _vision_cfg():
    base = _get_setting("xhs_vision_base", "").strip() or _llm_cfg()["base"]
    key = _get_setting("xhs_vision_key", "").strip() or _llm_cfg()["key"]
    model = _get_setting("xhs_vision_model", "").strip()
    if not model and "opencode.ai" in base:
        model = "deepseek-v4-flash-vision-exp"
    return {"base": base.rstrip("/"), "key": key, "model": model}


def _llm_vision(system, user, images, timeout=180):
    cfg = _vision_cfg()
    if not cfg["key"]:
        raise RuntimeError("未配置视觉模型 Key")
    if not cfg["model"]:
        raise RuntimeError("未配置视觉模型名（图片转写需要视觉模型）")
    content = [{"type": "text", "text": user}]
    for b64, mime in images:
        content.append({"type": "image_url", "image_url": {"url": "data:%s;base64,%s" % (mime, b64)}})
    r = requests.post(
        cfg["base"] + "/chat/completions",
        json={"model": cfg["model"], "messages": [{"role": "system", "content": system}, {"role": "user", "content": content}], "temperature": 0.1},
        headers=_llm_headers(cfg),
        timeout=timeout,
    )
    d = r.json()
    if r.status_code != 200 or "choices" not in d:
        raise RuntimeError("API错误: %s" % ((d.get("error") or {}).get("message") or r.text[:200]))
    return d["choices"][0]["message"]["content"]


def _transcribe_note_images(images, cookie="", proxy="", max_images=4):
    if not images:
        return ""
    try:
        import xhs
    except Exception:
        return ""
    from io import BytesIO
    datas = []
    for u in images[:max_images]:
        try:
            raw = xhs.download_image(u, cookie=cookie, proxy=proxy, timeout=20)
        except Exception:
            continue
        low = u.lower()
        if ".png" in low or "png" in low:
            mime = "image/png"
        elif ".webp" in low or "webp" in low:
            # 多数视觉接口不认 webp，统一转 png
            try:
                from PIL import Image
                img = Image.open(BytesIO(raw)).convert("RGB")
                out = BytesIO()
                img.save(out, format="PNG")
                raw = out.getvalue()
                mime = "image/png"
            except Exception:
                mime = "image/webp"
        else:
            mime = "image/jpeg"
        datas.append((base64.b64encode(raw).decode("ascii"), mime))
    if not datas:
        return ""
    system = "你是面经整理助手。请把图片中的面试经验/题目内容逐条转成文字，保留问题原文与关键信息，不要遗漏，不要编造。只输出转写正文。"
    try:
        return (_llm_vision(system, "请转写以下面试经验笔记图片：", datas) or "").strip()
    except Exception:
        return ""


def _playwright_installed():
    try:
        import playwright  # noqa
        return True
    except Exception:
        return False


@app.get("/api/xhs/config")
def xhs_config():
    return {
        "cookie": _get_setting("xhs_cookie", ""),
        "vision_model": _get_setting("xhs_vision_model", ""),
        "vision_base": _get_setting("xhs_vision_base", ""),
        "vision_key": _get_setting("xhs_vision_key", ""),
        "has_playwright": _playwright_installed(),
    }


@app.post("/api/xhs/config")
def xhs_set_config(body: dict):
    for k, v in body.items():
        if k in ("xhs_cookie", "xhs_vision_model", "xhs_vision_base", "xhs_vision_key"):
            _set_setting(k, str(v or "").strip())
    return {"ok": True}


@app.post("/api/xhs/test")
def xhs_test(body: dict):
    url = (body.get("url") or "").strip()
    if not url.startswith(("http://", "https://")):
        raise HTTPException(400, "URL 不合法")
    cookie = _get_setting("xhs_cookie", "").strip()
    proxy = _get_setting("jing_proxy", "").strip()
    note = None
    for p in ([None] + ([proxy] if proxy else [])):
        note = xhs.fetch_note(url, cookie=cookie, proxy=p or "")
        if note:
            break
    if not note:
        raise HTTPException(400, "抓取失败：直连与 Playwright 均未取到内容。可能需要登录 Cookie（设置-小红书-粘贴 Cookie）或平台反爬拦截。")
    result = {"title": note.get("title", ""), "desc": note.get("desc", ""), "images": len(note.get("images") or []), "reader": note.get("reader", "xhs")}
    if note.get("images") and not (note.get("desc") or "").strip():
        try:
            trans = _transcribe_note_images(note["images"], cookie=cookie, proxy=proxy)
            result["transcribed"] = trans[:4000]
        except Exception as e:
            result["transcribe_error"] = str(e)[:200]
    return result


@app.get("/api/jing/config")
def jing_config():
    return {"reader": _get_setting("jing_reader", "auto"), "proxy": _get_setting("jing_proxy", "")}


@app.post("/api/jing/config")
def jing_set_config(body: dict):
    if "reader" in body:
        v = str(body["reader"] or "").strip()
        if v not in ("auto", "direct", "jina"):
            raise HTTPException(400, "非法方式")
        _set_setting("jing_reader", v)
    if "proxy" in body:
        _set_setting("jing_proxy", str(body["proxy"] or "").strip())
    return {"ok": True}


_BAD_CONTENT_HINTS = (
    "页面不存在", "页面访问失败", "内容已删除", "该内容已被删除", "内容已被删除",
    "登录小红书", "请登录", "登录后", "验证码", "访问过于频繁",
    "Page Not Found", "404 Not Found",
)


def _looks_valid_content(text):
    t = (text or "").strip()
    if len(t) < 10:
        return False
    for h in _BAD_CONTENT_HINTS:
        if h in t:
            return False
    return True


def _fetch_url_text(url):
    """抓取任意链接正文，返回 (content_text, reader)。小红书走分层管线，其余走直连/Jina。"""
    cfg_proxy = _get_setting("jing_proxy", "").strip()
    if xhs.is_xhs_url(url):
        xhs_cookie = _get_setting("xhs_cookie", "").strip()
        note = None
        for proxy in ([None] + ([cfg_proxy] if cfg_proxy else [])):
            note = xhs.fetch_note(url, cookie=xhs_cookie, proxy=proxy or "")
            if note:
                break
        if note:
            parts = []
            if note.get("title"):
                parts.append("【标题】" + str(note["title"]))
            if note.get("desc"):
                parts.append("【正文】" + str(note["desc"]))
            images = note.get("images") or []
            if images and not (note.get("desc") or "").strip():
                try:
                    trans = _transcribe_note_images(images, cookie=xhs_cookie, proxy=cfg_proxy)
                    if trans:
                        parts.append("【图片转写】" + trans)
                except Exception:
                    pass
            content = "\n\n".join(parts).strip()
            if content:
                return content[:8000], note.get("reader", "xhs")
    headers = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120 Safari/537.36"}
    reader = _get_setting("jing_reader", "auto")
    methods = ["direct", "jina"] if reader == "auto" else ([reader] if reader in ("direct", "jina") else ["direct"])
    proxies = [None] + ([cfg_proxy] if cfg_proxy else [])
    for method in methods:
        for proxy in proxies:
            try:
                text = _fetch_direct(url, headers, proxy) if method == "direct" else _fetch_jina(url, proxy)
                if _looks_valid_content(text):
                    return text.strip()[:8000], method
            except Exception:
                continue
    return "", ""


@app.post("/api/jing/fetch")
def jing_fetch(body: dict):
    url = (body.get("url") or "").strip()
    if not url.startswith(("http://", "https://")):
        raise HTTPException(400, "URL 不合法")
    content, reader = _fetch_url_text(url)
    if not content:
        raise HTTPException(400, "抓取失败：平台反爬或网络受限。可在「设置」配置代理 / 切换 Jina / 填小红书 Cookie，或直接手动粘贴正文")
    return {"content": content, "reader": reader, "proxy": bool(_get_setting("jing_proxy", "").strip())}


@app.post("/api/jing/batch_import")
def jing_batch_import(body: dict):
    links = [str(x).strip() for x in (body.get("links") or []) if str(x).strip()]
    folder_id = int(body.get("folder_id") or 0)
    if not links:
        raise HTTPException(400, "请粘贴至少一条链接")
    if len(links) > 20:
        raise HTTPException(400, "一次最多 20 条链接")
    results = []
    for url in links:
        res = {"url": url}
        try:
            content, _ = _fetch_url_text(url)
            if not content:
                res.update({"status": "fail", "error": "未抓到内容"})
                results.append(res)
                continue
            d = _llm_json(JING_PARSE_PROMPT, content[:6000], kind="面经")
            company = str(d.get("company", "") or "").strip()
            title = str(d.get("title", "") or "").strip()
            questions = [str(q).strip() for q in (d.get("questions") or []) if str(q).strip()][:15]
            summary = str(d.get("summary", "") or "").strip()
            c = _conn()
            try:
                with c:
                    cur = c.execute(
                        "INSERT INTO jings (company,title,source_url,content,questions,summary,is_fav,folder_id,source,review_id,created_at) VALUES (?,?,?,?,?,?,?,?,?,?,?)",
                        (company or "未知公司", title or "面经", url, content, json.dumps(questions, ensure_ascii=False),
                         summary, 0, folder_id, "", 0, time.strftime("%Y-%m-%d %H:%M")),
                    )
                res.update({"status": "ok", "id": cur.lastrowid, "company": company, "title": title, "questions": len(questions)})
            finally:
                c.close()
        except HTTPException as e:
            res.update({"status": "fail", "error": str(e.detail)[:150]})
        except Exception as e:
            res.update({"status": "fail", "error": str(e)[:150]})
        results.append(res)
    return {"results": results}


def _company_fav_questions(company):
    if not (company or "").strip():
        return []
    c = _conn()
    try:
        rows = c.execute("SELECT questions FROM jings WHERE company=? AND is_fav=1", (company,)).fetchall()
    finally:
        c.close()
    qs = []
    for r in rows:
        try:
            qs += json.loads(r["questions"] or "[]")
        except Exception:
            pass
    return [str(q).strip() for q in qs if str(q).strip()][:20]


def _interview_q_context(company, learn=True):
    if not (company or "").strip():
        return []
    base = _company_fav_questions(company)
    if not learn:
        return base
    c = _conn()
    try:
        rows = c.execute("SELECT questions FROM jings WHERE company=? ORDER BY is_fav DESC, id DESC", (company,)).fetchall()
    finally:
        c.close()
    qs = list(base)
    for r in rows:
        try:
            qs += json.loads(r["questions"] or "[]")
        except Exception:
            pass
    seen, out = set(), []
    for q in qs:
        q = str(q).strip()
        if q and q not in seen:
            seen.add(q)
            out.append(q)
    return out[:30]


def _my_interview_jings():
    c = _conn()
    try:
        rows = c.execute("SELECT company,title,questions,summary FROM jings WHERE source='review' ORDER BY id DESC").fetchall()
    finally:
        c.close()
    out = []
    for r in rows:
        try:
            qs = json.loads(r["questions"] or "[]")
        except Exception:
            qs = []
        out.append({"company": r["company"], "title": r["title"],
                    "questions": [str(q) for q in qs], "summary": r["summary"] or ""})
    return out


STYLE_PROFILE_PROMPT = (
    "你是资深面试教练。以下是求职者记录的真实面试问题（公司·岗位·被问问题）。\n"
    "请提炼一份「真实面试风格画像」，用于指导 AI 模拟面试官模拟该求职者真实经历的面试风格。\n"
    "包含：1) 出题偏好（常问的题型/主题/方向）2) 难度特点 3) 追问方式 4) 该求职者的表现短板（哪些题容易答不好）\n"
    "用中文，150 字内，直接输出画像正文，不要多余解释。"
)


def _interview_style_profile():
    mine = _my_interview_jings()
    if not mine:
        return ""
    src = "\n".join("[%s %s]\n%s" % (m["company"] or "?", m["title"] or "?", "\n".join(m["questions"])) for m in mine)
    src = src[:6000]
    key = "style_prof:" + hashlib.sha256(src.encode("utf-8")).hexdigest()
    hit = _cache_get(key)
    if hit is not None:
        return hit
    try:
        prof = _llm_text(STYLE_PROFILE_PROMPT, src, kind="面试画像").strip()
        _cache_set(key, prof)
        return prof
    except HTTPException as e:
        log.warning("面试画像生成失败: %s", e)
        return ""


def _save_mock_jing(job, msgs):
    qs = []
    for m in msgs:
        if m.get("role") == "assistant":
            q = str(m.get("question", "") or "").strip()
            if q and "点击结束" not in q and "结束生成报告" not in q:
                qs.append(q)
    if not qs:
        return
    folder = _mock_folder()
    content = "\n".join("Q: " + q for q in qs)
    try:
        c = _conn()
        try:
            with c:
                c.execute(
                    "INSERT INTO jings (company,title,source_url,content,questions,summary,is_fav,folder_id,source,review_id,created_at) VALUES (?,?,?,?,?,?,?,?,?,?,?)",
                    (job["company"], (job["company"] or "?") + " · " + (job["title"] or "?") + " · 模拟面试", "",
                     content, json.dumps(qs, ensure_ascii=False), "AI 模拟面试问题", 0, folder, "mock", 0,
                     time.strftime("%Y-%m-%d %H:%M")),
                )
        finally:
            c.close()
    except Exception as e:
        log.warning("模拟面试问题入库失败: %s", e)


@app.get("/api/jing/folders")
def list_folders():
    c = _conn()
    try:
        rows = c.execute(
            "SELECT f.id, f.name, f.created_at, (SELECT COUNT(*) FROM jings j WHERE j.folder_id=f.id) AS cnt "
            "FROM jing_folders f ORDER BY f.id ASC"
        ).fetchall()
        return [dict(r) for r in rows]
    finally:
        c.close()


@app.post("/api/jing/folders")
def add_folder(body: dict):
    name = str(body.get("name") or "").strip()
    if not name:
        raise HTTPException(400, "文件夹名不能为空")
    c = _conn()
    try:
        with c:
            cur = c.execute("INSERT INTO jing_folders (name,created_at) VALUES (?,?)",
                            (name, time.strftime("%Y-%m-%d %H:%M")))
        return {"id": cur.lastrowid}
    finally:
        c.close()


@app.delete("/api/jing/folders/{fid}")
def del_folder(fid: int):
    c = _conn()
    try:
        with c:
            cur = c.execute("DELETE FROM jing_folders WHERE id=?", (fid,))
            c.execute("UPDATE jings SET folder_id=0 WHERE folder_id=?", (fid,))
        if cur.rowcount == 0:
            raise HTTPException(404, "文件夹不存在")
        return {"ok": True}
    finally:
        c.close()


@app.get("/api/monitor/summary")
def monitor_summary():
    c = _conn()
    try:
        usage = c.execute("SELECT kind, COUNT(*) AS n, SUM(in_chars) AS ic, SUM(out_chars) AS oc, AVG(latency) AS al, SUM(latency) AS tl FROM usage GROUP BY kind").fetchall()
        total = c.execute("SELECT COUNT(*) AS n, AVG(latency) AS al, SUM(latency) AS tl FROM usage").fetchone()
        recent = c.execute("SELECT * FROM op_log ORDER BY id DESC LIMIT 30").fetchall()
        err_total = c.execute("SELECT COUNT(*) AS n FROM op_log WHERE level='error'").fetchone()[0]
        cache = c.execute("SELECT COUNT(*) AS n FROM cache").fetchone()[0]
        jobs_n = c.execute("SELECT COUNT(*) AS n FROM jobs").fetchone()[0]
        jings_n = c.execute("SELECT COUNT(*) AS n FROM jings").fetchone()[0]
        folders_n = c.execute("SELECT COUNT(*) AS n FROM jing_folders").fetchone()[0]
    finally:
        c.close()
    db_size = os.path.getsize(DB_PATH) if os.path.exists(DB_PATH) else 0
    cfg = _llm_cfg()
    feishu_ok = False
    try:
        feishu_ok = bool(feishu.get_token())
    except Exception:
        pass
    llm_ok, llm_err = False, ""
    try:
        if cfg["key"]:
            _llm_text("回复OK", "hi", override=cfg, use_cache=False)
            llm_ok = True
        else:
            llm_err = "未配置 API Key"
    except Exception as e:
        llm_err = str(e)[:150]
    return {
        "modules": [{"kind": r["kind"], "calls": r["n"], "in_chars": r["ic"] or 0, "out_chars": r["oc"] or 0,
                     "avg_latency": round(r["al"] or 0, 2), "total_latency": round(r["tl"] or 0, 2)} for r in usage],
        "total_calls": total["n"] or 0,
        "total_latency": round(total["tl"] or 0, 2),
        "avg_latency": round(total["al"] or 0, 2),
        "error_count": err_total,
        "recent_log": [dict(r) for r in recent],
        "cache_count": cache,
        "jobs_count": jobs_n,
        "jings_count": jings_n,
        "folders_count": folders_n,
        "db_size": db_size,
        "llm_ok": llm_ok, "llm_err": llm_err,
        "model": cfg["model"],
        "feishu_ok": feishu_ok,
        "feishu_configured": bool(ENV("FEISHU_APP_ID") and ENV("FEISHU_APP_SECRET")),
    }


AVATAR_DIR = ROOT / "static" / "avatars"


@app.post("/api/avatar/upload")
def avatar_upload(file: UploadFile = File(...)):
    name = (file.filename or "").lower()
    ext = name.rsplit(".", 1)[-1] if "." in name else ""
    if ext not in ("jpg", "jpeg", "png", "webp"):
        raise HTTPException(400, "仅支持 jpg/jpeg/png/webp")
    data = file.file.read()
    if len(data) > 5 * 1024 * 1024:
        raise HTTPException(400, "图片过大（限 5MB）")
    AVATAR_DIR.mkdir(parents=True, exist_ok=True)
    fname = "av_%s.%s" % (hashlib.sha1(data).hexdigest()[:12], ext)
    (AVATAR_DIR / fname).write_bytes(data)
    return {"url": "/avatars/" + fname}


@app.get("/avatars/{fname}")
def avatar_static(fname: str):
    if "/" in fname or "\\" in fname or not fname:
        raise HTTPException(400, "非法文件名")
    p = AVATAR_DIR / fname
    if not p.exists():
        raise HTTPException(404, "头像不存在")
    return FileResponse(p)


PHOTO_PATH = ROOT / "static" / "photo.jpg"


@app.post("/api/photo/upload")
def photo_upload(file: UploadFile = File(...)):
    name = (file.filename or "").lower()
    ext = name.rsplit(".", 1)[-1] if "." in name else ""
    if ext not in ("jpg", "jpeg", "png", "webp"):
        raise HTTPException(400, "仅支持 jpg/jpeg/png/webp")
    data = file.file.read()
    if len(data) > 5 * 1024 * 1024:
        raise HTTPException(400, "图片过大（限 5MB）")
    try:
        from io import BytesIO
        from PIL import Image
        img = Image.open(BytesIO(data)).convert("RGB")
        out = BytesIO()
        img.save(out, format="JPEG", quality=92)
        data = out.getvalue()
    except Exception:
        pass
    PHOTO_PATH.write_bytes(data)
    return {"ok": True}


@app.get("/api/photo/info")
def photo_info():
    return {"has": PHOTO_PATH.exists()}


@app.get("/api/photo")
def photo_serve():
    if not PHOTO_PATH.exists():
        raise HTTPException(404, "未上传证件照")
    return FileResponse(PHOTO_PATH)


@app.get("/")
def home():
    return FileResponse(ROOT / "static" / "index.html")
